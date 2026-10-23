"""Domain-agnostic FAST-workbook builder.

Emits a FAST-structured ``.xlsx`` that *calculates* — live formulas, not lifted totals —
in the house style (see :mod:`omd.sheet.formatting`). The FAST discipline it enforces:

- **Inputs / source tables / Calcs / Statement are separate sheets** (one thing per sheet).
- **Time flows left -> right**: the same period columns on every sheet.
- **One row = one calculation, one formula shifted across the row** — you supply a single
  per-period formula function and the builder copies it across time.
- **Transparency**: inputs are shaded; every statement number links back to a Calcs cell.

The builder is data-free: a *spec* (Python module) drives it with the actual lines and formulas.
Thread known ``values`` alongside each formula so formula cells also get the smart-number
format (whole -> no decimals, fractional -> 2dp) instead of a blanket 2dp fallback.

Sketch::

    wb = Workbook("My model", periods=["Jan", "Feb", "Mar"], opening=True)
    ins = wb.inputs("Inputs", {"cash0": ("Opening cash", 1000.0)})
    led = wb.table("Ledger", ["month", "kind", "amount"], rows)
    calc = wb.calcs()
    calc.section("Flows")
    calc.row("rev", "  Revenue", lambda c: c.sumifs(led.col("amount"),
             f'{led.col("kind")},"sale"', f'{led.col("month")},{c.critcell}'), values=rev_vals)
    calc.roll("cash", "  Cash", opening=f"={ins['cash0']}", add=["rev"], values=cash_vals)
    wb.statement("Accounts", [("P&L", [("Revenue", "rev")])], title="My Model")
    wb.autosize(); wb.save(path)
"""

from __future__ import annotations

from collections.abc import Callable, Mapping, Sequence
from datetime import datetime
from typing import Any

from openpyxl import Workbook as _OpenpyxlWorkbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from omd.style import StyleProfile, get_profile

from . import formatting as F

_EPOCH = datetime(2026, 1, 1)  # fixed doc timestamp (no wall-clock) -> deterministic content


# ==============================================================================================
# Source data table (a SUMIFS / LOOKUP source)
# ==============================================================================================
class Table:
    """A flat source-data sheet. Exposes absolute column ranges for SUMIFS / LOOKUP."""

    def __init__(
        self, ws: Worksheet, name: str, headers: Sequence[str], first_row: int, last_row: int
    ):
        self.ws = ws
        self.name = name
        self._first = first_row
        self._last = last_row
        self._col_of = {h: get_column_letter(i + 1) for i, h in enumerate(headers)}

    def col(self, header: str) -> str:
        """Absolute range for a named column, e.g. ``'Ledger'!$C$3:$C$120`` (for SUMIFS)."""
        c = self._col_of[header]
        return f"'{self.name}'!${c}${self._first}:${c}${self._last}"


# ==============================================================================================
# Calcs — one row per line, live formulas, time across the columns
# ==============================================================================================
class _Ctx:
    """Per-column context handed to each row's formula function (FAST: one formula, shifted)."""

    def __init__(self, calcs: Calcs, col: str, prev_col: str | None, period_index: int):
        self._c = calcs
        self.col = col
        self.prev_col = prev_col
        self.i = period_index  # 1-based period index

    @property
    def critcell(self) -> str:
        """The criterion cell for this column (the period header, row 2) — for SUMIFS."""
        return f"{self.col}${self._c._HDR_ROW}"

    @property
    def crit(self) -> Any:
        """The literal criterion value for this column (if a literal is preferred over a cell)."""
        return self._c._periods[self.i - 1]

    def here(self, key: str) -> str:
        """Ref to another Calcs line's cell in *this* column."""
        return f"{self.col}{self._c._rows[key]}"

    def prevref(self, key: str) -> str:
        """Ref to another Calcs line's cell in the *previous* column (roll-forward)."""
        if self.prev_col is None:
            raise ValueError("no previous column for the first period (need opening=True)")
        return f"{self.prev_col}{self._c._rows[key]}"

    def input(self, name: str) -> str:
        return self._c._wb._refs[name]

    @staticmethod
    def sumifs(value_range: str, *criteria: str) -> str:
        return f"=SUMIFS({value_range}," + ",".join(criteria) + ")"


class Calcs:
    """The working sheet: one row per line, every cell a live formula shifted across periods."""

    def __init__(self, wb: Workbook, ws: Worksheet, name: str):
        self._wb = wb
        self.ws = ws
        self.name = name
        self._rows: dict[str, int] = {}
        self._periods = wb._periods
        S = wb.style
        S.white(ws)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(wb._all_cols) + 1)
        S.title(ws.cell(row=1, column=1, value=f"{name} — one row per line; live formulas"))

        self._HDR_ROW = 2
        S.colhdr(ws.cell(row=2, column=1, value=""), right=False)
        if wb._open_col:
            S.colhdr(ws.cell(row=2, column=wb._open_col_idx, value="Open"))
        for label, letter in zip(wb._periods, wb._period_cols):
            cell = ws[f"{letter}2"]
            cell.value = label
            S.colhdr(cell)
        ws.column_dimensions["A"].width = 42
        self._row = 4  # first content row (leave row 3 as a spacer / driver band)
        self._timeline()

    def _timeline(self) -> None:
        """The period-position rows, in the engine core: a flag for the first period, a flag for the last, and the count
        of periods after this one. A row that behaves differently at an edge (or looks ahead) multiplies by one of
        these, gated, so no formula writes a different shape in one column and none reads a cell past the last period
        without a gate that switches it off."""
        wb, n = self._wb, len(self._wb._period_cols)
        first, last = wb._period_cols[0], wb._period_cols[-1]
        self.section("Timeline")
        self.row("first_period_flag", "  First period (flag)",
                 lambda x: f"=(COLUMN()=COLUMN(${first}${self._HDR_ROW}))*1", values=[1.0] + [0.0] * (n - 1))  # fmt: skip
        self.row("last_period_flag", "  Last period (flag)",
                 lambda x: f"=(COLUMN()=COLUMN(${last}${self._HDR_ROW}))*1", values=[0.0] * (n - 1) + [1.0])  # fmt: skip
        self.row("periods_left", "  Periods after this one (count)",
                 lambda x: f"=COLUMN(${last}${self._HDR_ROW})-COLUMN()", values=[float(n - 1 - i) for i in range(n)])  # fmt: skip
        self.skip_rows()

    # -- referencing --------------------------------------------------------------------------
    def rows(self) -> dict[str, int]:
        return dict(self._rows)

    @property
    def workbook(self) -> Workbook:
        """The :class:`Workbook` this sheet belongs to."""
        return self._wb

    @property
    def next_row(self) -> int:
        """The row the next ``row()``/``section()`` will write. Read it to predict a later row."""
        return self._row

    def skip_rows(self, n: int = 1) -> None:
        """Leave ``n`` blank rows (breathing room, or space reserved to fill later)."""
        if n < 0:
            raise ValueError("skip_rows needs n >= 0")
        self._row += n

    @property
    def header_row(self) -> int:
        """The row that holds the period headers."""
        return self._HDR_ROW

    # -- write a section header ---------------------------------------------------------------
    def section(self, label: str) -> None:
        self._wb.style.section(self.ws.cell(row=self._row, column=1, value=label))
        self._row += 1

    # -- write a line -------------------------------------------------------------------------
    def row(
        self,
        key: str,
        label: str,
        formula: Callable[[_Ctx], str],
        *,
        values: Sequence[float] | None = None,
        open_formula: str | None = None,
        open_value: float | None = None,
        open_known: float | None = None,
        fmt: str = F.NUM,
    ) -> str:
        """Write one line. ``formula(c)`` returns the per-period formula string.

        ``values`` (known numbers, one per period) drives the smart-number format on each
        formula cell — the fix for the blanket-2dp fallback. ``fmt`` overrides (e.g. ``RATE``).
        Returns ``key`` for chaining.
        """
        wb = self._wb
        rr = self._row
        self.ws.cell(row=rr, column=1, value=label)
        self._rows[key] = rr  # set before the loop so roll-forward can reference self

        # opening cell
        if wb._open_col and (open_formula is not None or open_value is not None):
            oc = self.ws[f"{wb._open_col}{rr}"]
            oc.value = open_formula if open_formula is not None else open_value
            base = open_value if open_value is not None else open_known
            oc.number_format = (
                self._wb.style.fmt_for(base, fmt)
                if base is not None
                else self._wb.style.fallback(fmt)
            )

        prev_col = wb._open_col
        for i, col in enumerate(wb._period_cols, start=1):
            ctx = _Ctx(self, col, prev_col, i)
            cell = self.ws[f"{col}{rr}"]
            cell.value = formula(ctx)
            known = values[i - 1] if values is not None else None
            cell.number_format = (
                self._wb.style.fmt_for(known, fmt)
                if known is not None
                else self._wb.style.fallback(fmt)
            )
            prev_col = col
        self._row += 1
        return key

    # -- roll-forward convenience -------------------------------------------------------------
    def roll(
        self,
        key: str,
        label: str,
        *,
        opening: str | float,
        add: Sequence[str] = (),
        sub: Sequence[str] = (),
        values: Sequence[float] | None = None,
        open_known: float | None = None,
        fmt: str = F.NUM,
    ) -> str:
        """Roll-forward line: each period = previous period + Σadd(this col) − Σsub(this col).

        ``opening`` is the Opening-column contents: a number, or a formula string (``=...``,
        e.g. a link to an Inputs cell). ``add``/``sub`` are other Calcs line keys.
        """
        if not self._wb._open_col:
            raise ValueError("roll() needs a workbook built with opening=True")

        def _fn(c: _Ctx) -> str:
            terms = [c.prevref(key)]
            terms += [f"+{c.here(k)}" for k in add]
            terms += [f"-{c.here(k)}" for k in sub]
            expr = terms[0] + "".join(terms[1:])
            return f"={expr}"

        is_formula = isinstance(opening, str)
        return self.row(
            key,
            label,
            _fn,
            values=values,
            open_formula=opening if is_formula else None,
            open_value=None if is_formula else float(opening),
            open_known=open_known if is_formula else None,
            fmt=fmt,
        )


# ==============================================================================================
# Workbook
# ==============================================================================================
class Workbook:
    """A FAST workbook under construction. See the module docstring for the flow."""

    def __init__(
        self,
        title: str,
        periods: Sequence[Any],
        *,
        opening: bool = False,
        subtitle: str | None = None,
        style: str | StyleProfile | None = None,
    ):
        if not periods:
            raise ValueError("at least one period column is required")
        self.title = title
        self.subtitle = subtitle
        self.style: StyleProfile = get_profile(style)
        self._periods = list(periods)
        self._wb = _OpenpyxlWorkbook()
        self._wb.remove(self._wb.active)  # start empty

        # column layout: A = labels; optional Open col; then one column per period
        self._open_col_idx = 2 if opening else None
        self._open_col = get_column_letter(2) if opening else None
        first_period_idx = 3 if opening else 2
        self._period_cols = [get_column_letter(first_period_idx + i) for i in range(len(periods))]
        self._all_cols = ([self._open_col] if self._open_col else []) + self._period_cols

        self._refs: dict[str, str] = {}
        self._inputs_ws: Worksheet | None = None
        self._input_row = 4
        self._calcs: Calcs | None = None

    # -- columns ------------------------------------------------------------------------------
    @property
    def period_cols(self) -> tuple[str, ...]:
        """Column letters of the periods, in time order."""
        return tuple(self._period_cols)

    @property
    def open_col(self) -> str | None:
        """Column letter of the Opening column, or None if the workbook has none."""
        return self._open_col

    @property
    def all_cols(self) -> tuple[str, ...]:
        """The Opening column (if any) followed by the period columns."""
        return tuple(self._all_cols)

    def set_layout(self, *, open_col: int | None, first_period_col: int) -> None:
        """Move the Opening and period columns (1-based indexes), e.g. to make room for a units
        column. Call it before :meth:`calcs`; a Calcs sheet already written would not move."""
        if self._calcs is not None:
            raise ValueError("set_layout must be called before calcs()")
        if open_col is not None and open_col >= first_period_col:
            raise ValueError("the Opening column must sit left of the first period column")
        self._open_col_idx = open_col
        self._open_col = get_column_letter(open_col) if open_col else None
        self._period_cols = [
            get_column_letter(first_period_col + i) for i in range(len(self._periods))
        ]
        self._all_cols = ([self._open_col] if self._open_col else []) + self._period_cols

    # -- sheets -------------------------------------------------------------------------------
    @property
    def inputs_sheet(self) -> Worksheet | None:
        """The Inputs worksheet, or None before :meth:`inputs` has been called."""
        return self._inputs_ws

    @property
    def sheet_names(self) -> list[str]:
        return list(self._wb.sheetnames)

    def sheet(self, name: str) -> Worksheet:
        """A worksheet by name."""
        return self._wb[name]

    def add_sheet(self, title: str, index: int | None = None) -> Worksheet:
        """Add an empty worksheet, at the end or at ``index``."""
        return self._wb.create_sheet(title, index)

    def rename_sheet(self, old: str, new: str) -> None:
        """Change a sheet's title only. Formulas that name the old sheet are NOT rewritten; the
        caller must rewrite them (and `Calcs.name` if it was Calcs)."""
        self._wb[old].title = new

    def order_sheets(self, names: Sequence[str]) -> None:
        """Put the sheets in this order. ``names`` must name every sheet exactly once."""
        if sorted(names) != sorted(self._wb.sheetnames):
            raise ValueError(
                f"order_sheets needs every sheet once: have {sorted(self._wb.sheetnames)}, "
                f"got {sorted(names)}"
            )
        self._wb._sheets = [self._wb[n] for n in names]

    @property
    def full_calc_on_load(self) -> bool:
        return bool(self._wb.calculation.fullCalcOnLoad)

    @full_calc_on_load.setter
    def full_calc_on_load(self, value: bool) -> None:
        self._wb.calculation.fullCalcOnLoad = bool(value)

    # -- Inputs -------------------------------------------------------------------------------
    def inputs(self, section: str, entries: Mapping[str, Any]) -> dict[str, str]:
        """Add shaded input cells under a section header; return ``{name: cell-ref}``.

        ``entries`` maps a name to either a scalar value, or a ``(label, value)`` /
        ``(label, value, fmt)`` tuple. Names accumulate across calls into a single Inputs sheet
        and are referable from formulas via the returned refs (and ``ctx.input(name)``).
        """
        ws = self._inputs_ws
        if ws is None:
            ws = self._wb.create_sheet("Inputs")
            self.style.white(ws)
            ws.merge_cells("A1:C1")
            self.style.title(ws.cell(row=1, column=1, value=self.title))
            ws.merge_cells("A2:C2")
            self.style.subtitle(
                ws.cell(
                    row=2,
                    column=1,
                    value="Shaded cells are hardcoded inputs — edit to recalculate.",
                )
            )
            self._inputs_ws = ws

        self.style.section(ws.cell(row=self._input_row, column=1, value=section))
        self._input_row += 1
        out: dict[str, str] = {}
        for name, spec in entries.items():
            label, value, fmt = self._parse_input(name, spec)
            r = self._input_row
            ws.cell(row=r, column=1, value=label)
            cell = ws.cell(row=r, column=2, value=value)
            self.style.mark_input(cell, self.style.fmt_for(value, fmt))
            ref = f"Inputs!$B${r}"
            self._refs[name] = ref
            out[name] = ref
            self._input_row += 1
        self._input_row += 1  # gap before the next section
        return out

    @staticmethod
    def _parse_input(name: str, spec: Any) -> tuple[str, Any, str]:
        if isinstance(spec, tuple):
            if len(spec) == 3:
                return spec[0], spec[1], spec[2]
            if len(spec) == 2:
                return spec[0], spec[1], F.NUM
            raise ValueError(f"input {name!r}: expected value or (label, value[, fmt])")
        return name, spec, F.NUM

    # -- source table -------------------------------------------------------------------------
    def table(
        self,
        name: str,
        headers: Sequence[str],
        rows: Sequence[Sequence[Any] | Mapping[str, Any]],
        *,
        title: str | None = None,
    ) -> Table:
        """Write a flat source-data sheet (SUMIFS/LOOKUP source). Numeric cells get smart format."""
        ws = self._wb.create_sheet(name)
        self.style.white(ws)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        self.style.title(ws.cell(row=1, column=1, value=title or name))
        for j, h in enumerate(headers, start=1):
            self.style.colhdr(ws.cell(row=2, column=j, value=h), right=False)
        first = 3
        r = first
        for record in rows:
            if isinstance(record, Mapping):
                record = [record.get(h) for h in headers]
            for j, v in enumerate(record, start=1):
                cell = ws.cell(row=r, column=j, value=v)
                if isinstance(v, bool):
                    cell.number_format = self.style.acct0
                elif isinstance(v, (int, float)):
                    cell.number_format = self.style.smart_num(v)
            r += 1
        last = r - 1
        ws.freeze_panes = "A3"
        return Table(ws, name, headers, first, last)

    # -- Calcs --------------------------------------------------------------------------------
    def calcs(self, name: str = "Calcs") -> Calcs:
        ws = self._wb.create_sheet(name)
        self._calcs = Calcs(self, ws, name)
        return self._calcs

    # -- Statement (reviewer-facing presentation, links to Calcs) -----------------------------
    def statement(
        self,
        name: str,
        blocks: Sequence[tuple[str, Sequence[tuple]]],
        *,
        title: str | None = None,
        subtitle: str | None = None,
        calcs: Calcs | None = None,
    ) -> None:
        """IB-style presentation sheet. ``blocks`` = ``[(section_title, items), ...]`` where each
        item is ``(label, calcs_key)`` plus optional flags ``{"total": True}`` / ``{"grand": True}``
        or ``(label, None)`` for a blank spacer row. Every value links to the Calcs cell.
        """
        calcs = calcs or self._calcs
        if calcs is None:
            raise ValueError("statement() needs a Calcs sheet — call .calcs() first")
        R = calcs.rows()
        ws = self._wb.create_sheet(name)
        self.style.white(ws)
        ws.column_dimensions["A"].width = 3  # left margin

        n = len(self._periods)
        last_col = 3 + n - 1
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
        self.style.bold(
            ws.cell(row=2, column=2, value=title or self.title), size=14, color=self.style.accent
        )
        if subtitle or self.subtitle:
            ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last_col)
            self.style.subtitle(ws.cell(row=3, column=2, value=subtitle or self.subtitle))

        hdr = 5
        self.style.colhdr(ws.cell(row=hdr, column=2, value=""), right=False)
        for j, lab in enumerate(self._periods):
            self.style.colhdr(ws.cell(row=hdr, column=3 + j, value=str(lab)))

        acols = self._period_cols
        row = hdr + 1
        first_section = True
        for section_title, items in blocks:
            if not first_section:
                row += 2
            first_section = False
            self.style.section(ws.cell(row=row, column=2, value=section_title), size=12)
            for col in range(2, last_col + 1):
                ws.cell(row=row, column=col).border = self.style.hdr_border
            row += 1
            for item in items:
                label, key = item[0], item[1]
                opts = item[2] if len(item) > 2 else {}
                if key is None:  # blank spacer row
                    row += 1
                    continue
                total, grand = opts.get("total", False), opts.get("grand", False)
                lc = ws.cell(row=row, column=2, value=label)
                if total or grand:
                    self.style.bold(lc)
                    lc.border = self.style.total_top if total else self.style.total_bottom
                for j, ccol in enumerate(acols):
                    cell = ws.cell(row=row, column=3 + j, value=f"={calcs.name}!{ccol}{R[key]}")
                    cell.number_format = self.style.acct0  # statements: whole numbers
                    cell.alignment = F.RIGHT
                    if total or grand:
                        self.style.bold(cell)
                        cell.border = self.style.total_top if total else self.style.total_bottom
                row += 1
        ws.freeze_panes = "C6"

    # -- finalise -----------------------------------------------------------------------------
    def autosize(self) -> None:
        for ws in self._wb.worksheets:
            F.autosize(ws)

    def save(self, path: str) -> None:
        self._wb.properties.created = _EPOCH
        self._wb.properties.modified = _EPOCH
        self._wb.save(str(path))

    def cells(self) -> dict[tuple[str, str], Any]:
        """All (sheet, coord) -> value pairs — for a determinism check."""
        return {
            (ws.title, c.coordinate): c.value
            for ws in self._wb.worksheets
            for row in ws.iter_rows()
            for c in row
        }
