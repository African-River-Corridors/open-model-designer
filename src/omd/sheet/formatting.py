"""House spreadsheet formatting kit — the house style, domain-agnostic.

Lifted (generalised) from a FAST workbook writer. Pure presentation:

- ``smart_num`` — thousands separators; decimals **only where a value genuinely has
  cents**; **never a trailing dot**.
- shaded hardcoded-input cells (light blue) so a reviewer sees what is editable;
- ruled headers/totals (navy rules, **no fills**), all-white background (gridlines off);
- ``autosize`` — size each column to its widest displayed cell (formulas costed via
  their number format);
- callers stamp a **fixed doc timestamp** for byte-stable/deterministic output.

No project data lives here — this module only knows about colours, borders and numbers.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# ---- number formats --------------------------------------------------------------------------
ACCT0 = "#,##0;(#,##0)"  # whole numbers (no decimal point at all)
ACCT2 = "#,##0.00;(#,##0.00)"  # two decimals (only where a value genuinely has cents)
RATE = "0.######"  # a rate / multiplier
DATEF = "yyyy-mm-dd"
NUM = "__smart__"  # sentinel: pick the decimals per value (see smart_num)
PRICE = "__price__"  # a price: always 2 decimals
MONEY = NUM  # back-compat alias


def smart_num(v: float) -> str:
    """Decimals only where a value needs them (see :class:`omd.style.StyleProfile`)."""
    from omd.style import DEFAULT

    return DEFAULT.smart_num(v)


def fmt_for(value, fmt: str = NUM) -> str:
    """Resolve a format token against a known value in the default profile."""
    from omd.style import DEFAULT

    return DEFAULT.fmt_for(value, fmt)


# ---- palette / borders (all-white bg; inputs shaded; headers/totals ruled, not filled) -------
NAVY = "1F2A44"
MUTED = "55617A"
INPUT_FILL = PatternFill("solid", fgColor="E8F0FE")  # light blue = hardcoded, editable input
_THIN = Side(style="thin", color="8894AC")
_RULE = Side(style="thin", color=NAVY)
TOP = Border(top=_THIN)
TOP_DBL = Border(top=Side(style="double", color=NAVY))
HDR_BORDER = Border(bottom=_RULE)
RIGHT = Alignment(horizontal="right")


def bold(cell, size=None, color=None) -> None:
    cell.font = Font(bold=True, size=size, color=color)


def title(cell, size: int = 12) -> None:
    """Sheet title: bold navy, slightly larger."""
    cell.font = Font(bold=True, size=size, color=NAVY)


def subtitle(cell) -> None:
    cell.font = Font(italic=True, color=MUTED)


def colhdr(cell, right: bool = True) -> None:
    """Column/table header: bold navy text with a bottom rule (no fill)."""
    cell.font = Font(bold=True, color=NAVY)
    cell.border = HDR_BORDER
    if right:
        cell.alignment = RIGHT


def mark_input(cell, fmt: str | None = None) -> None:
    """Shade a hardcoded input cell so it reads as editable."""
    cell.fill = INPUT_FILL
    if fmt:
        cell.number_format = fmt


def white(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False


# ---- auto-size -------------------------------------------------------------------------------
def _est_width(cell) -> int:
    """Estimate the displayed width of a cell (handles formulas via their number format)."""
    v = cell.value
    if v is None or v == "":
        return 0
    fmt = cell.number_format or ""
    is_formula = isinstance(v, str) and v.startswith("=")
    if "yyyy" in fmt:
        return len("2026-06-30")
    if fmt.startswith("0."):  # a rate
        return len("99.999999") if is_formula else len(f"{v:.6f}".rstrip("0").rstrip("."))
    if "#,##0" in fmt:  # a grouped number
        return len("(999,999,999)") if is_formula else len(f"{v:,.0f}") + (2 if v < 0 else 0)
    if is_formula:
        return 12
    return len(str(v))


def lint_widths(
    ws: Worksheet,
    *,
    value_cols,
    note_col=None,
    value_min: int = 14,
    value_max: int = 24,
    note_max: int = 120,
    pad: int = 2,
    skip_rows=(),
) -> list[str]:
    """Make a presentation sheet's columns behave, and say what it had to do.

    ``autosize`` sizes every column to its own widest cell, which is right for a
    calculation sheet and wrong for a presentation one: three columns holding the same
    KIND of number end up three different widths because one of them happens to carry a
    longer stray string, and a notes column ends up wrapping onto three lines.

    Two rules, both of them things a reader notices immediately:

    * **Every value column gets the SAME width** — the widest any of them needs, clamped.
      Columns that hold the same kind of number should look like each other.
    * **The notes column is sized not to wrap**, up to ``note_max``. Past that it keeps
      wrapping, and the note is reported so it can be shortened at source instead.

    Returns the findings. It is a linter: it fixes what it safely can and TELLS YOU about
    what it cannot, rather than silently leaving a sheet that looks wrong.
    """
    findings: list[str] = []
    cols = [c.upper() for c in value_cols]
    # A footer (provenance, a signature block) is not a table and its columns mean nothing —
    # pass its rows here rather than have the linter read prose as a mis-placed note.
    skip = set(skip_rows)

    widest = value_min
    for col in cols:
        for cell in ws[col]:
            if cell.row in skip:
                continue
            w = _est_width(cell)
            # A stray long string in a value column is a MISTAKE, not a width problem —
            # it means a note was written into the wrong column. Say so; do not widen.
            if w > value_max and isinstance(cell.value, str) and not cell.value.startswith("="):
                findings.append(
                    f"{ws.title}!{cell.coordinate}: {w}-char text in value column {col} — "
                    f"a note in the wrong column? {cell.value[:48]!r}"
                )
                continue
            widest = max(widest, w + pad)
    widest = min(widest, value_max)
    for col in cols:
        ws.column_dimensions[col].width = widest

    if note_col:
        nc = note_col.upper()
        longest, worst = 0, None
        for cell in ws[nc]:
            if cell.row in skip:
                continue
            v = cell.value
            if not isinstance(v, str) or v.startswith("=") or not v.strip():
                continue
            if len(v) > longest:
                longest, worst = len(v), cell
        ws.column_dimensions[nc].width = min(longest + pad, note_max)
        if longest + pad > note_max and worst is not None:
            findings.append(
                f"{ws.title}!{worst.coordinate}: note is {longest} chars and will wrap at "
                f"{note_max} — shorten it at source: {worst.value[:48]!r}"
            )
    return findings


def autosize(ws: Worksheet, *, min_w: int = 6, max_w: int = 52, pad: int = 2) -> None:
    """Size each content column to fit its widest cell. Merged cells (titles) are ignored."""
    merged = {
        coord
        for mr in ws.merged_cells.ranges
        for row in ws[mr.coord]
        for coord in (c.coordinate for c in row)
    }
    widths: dict[str, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.coordinate in merged:
                continue
            w = _est_width(cell)
            if w:
                widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), w)
    for col, w in widths.items():
        ws.column_dimensions[col].width = min(max(w + pad, min_w), max_w)
