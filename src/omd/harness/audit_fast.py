"""Audit a generated workbook against the mechanically checkable FAST Standard 02c rules.

FAST Standard 02c (July 2019) is CC-BY 4.0: https://fast-standard.org. Rule numbers below are
quoted from Appendix C, "Rules in Short". This checks the rules a script can check; the rest
(labelling quality, sheet chaptering, presentation) still need a human read.

`declared` maps (rule code, row label) to the reason a breach is accepted. Declaring is not
hiding: declared breaches still print, marked as declared. A model passes its own declarations.
"""

from __future__ import annotations

import collections
import re
from dataclasses import dataclass, field

import openpyxl

CELLREF = re.compile(r"\$?[A-Z]{1,2}\$?\d+")
NUMLIT = re.compile(r"(?<![A-Z$\d.])(\d+(?:\.\d+)?)(?![\d.]*[)]?\s*[A-Z]*\$?\d)")
FUNC = re.compile(r"\b([A-Z][A-Z0-9.]+)\(")
OPS = re.compile(r"[+\-*/^]")

# 3.04-01.1 "Except when constants are universal" — months, hours, percent, sign, halving.
UNIVERSAL = {"0", "1", "2", "12", "100", "1000", "8760", "0.5", "-1"}


@dataclass
class FastAudit:
    findings: dict[str, list[tuple[int, str, str]]] = field(default_factory=dict)
    undeclared: int = 0
    n_rows: int = 0
    n_periods: int = 0
    duplicate_labels: dict[str, int] = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return self.undeclared == 0


def norm(f, col):
    """Strip the column letter so two columns' formulas can be compared structurally."""
    return re.sub(r"\$?([A-Z]{1,2})\$?(\d+)", lambda m: f"C{m.group(2)}", f)


def audit_fast(
    book: str,
    *,
    calcs_sheet: str = "Calcs",
    first_col: int = 2,
    declared: dict[tuple[str, str], str] | None = None,
    out=print,
) -> FastAudit:
    declared = declared or {}
    wb = openpyxl.load_workbook(book)
    ws = wb[calcs_sheet]
    nper = sum(1 for c in range(first_col, ws.max_column + 1) if ws.cell(row=2, column=c).value)

    rows = []
    for r in range(4, ws.max_row + 1):
        lab = ws.cell(row=r, column=1).value
        cells = [ws.cell(row=r, column=first_col + i).value for i in range(nper)]
        fs = [v for v in cells if isinstance(v, str) and v.startswith("=")]
        if not fs:
            continue
        rows.append((r, str(lab or "").strip(), cells, fs))

    findings = collections.defaultdict(list)
    for r, lab, cells, fs in rows:
        body = max(fs, key=len)[1:]

        # --- FAST 3.02-01 Formulas must be consistent (across the row) --------------------
        shapes = {
            norm(v, i) for i, v in enumerate(cells) if isinstance(v, str) and v.startswith("=")
        }
        if len(shapes) > 1:
            findings["3.02-01 formulas inconsistent across the row"].append(
                (r, lab, f"{len(shapes)} distinct formula shapes in {len(fs)} columns")
            )

        # --- FAST 3.03-01 Do not write a formula longer than your thumb -------------------
        if len(body) > 80:
            findings["3.03-01 formula longer than your thumb (>80 chars)"].append(
                (r, lab, f"{len(body)} chars")
            )

        # --- FAST 3.03-05 Use flags to limit use of IF; 3.03-07 never nest IFs ------------
        nif = len(re.findall(r"\bIF\(", body))
        if nif >= 2:
            findings["3.03-05/07 multiple or nested IFs — use a flag row"].append(
                (r, lab, f"{nif} IF()")
            )
        elif nif == 1:
            findings["3.03-05 IF() where a timing flag would do"].append((r, lab, ""))

        # --- FAST 3.04-01 Do not write formulas with embedded constants -------------------
        lits = [
            x
            for x in NUMLIT.findall(re.sub(r"\$?[A-Z]{1,2}\$?\d+", " ", body))
            if x not in UNIVERSAL
        ]
        if lits:
            findings["3.04-01 embedded constants"].append(
                (r, lab, ", ".join(sorted(set(lits))[:5]))
            )

        # --- FAST 4.01-02 Do not use the NPV function - ever ------------------------------
        used = set(FUNC.findall(body))
        if "NPV" in used:
            findings["4.01-02 NPV function used (never allowed)"].append((r, lab, ""))

        # --- FAST 3.05-04/09 Include a units designator in the label ----------------------
        if lab and not re.search(r"\(.*\)|\$/|%|\bx\b|—\s*\$", lab):
            findings["3.05-04/09 no unit designator in the label"].append((r, lab, ""))

    # --- FAST 3.05-03 All line items must have a unique label -----------------------------
    seen = collections.Counter(lab for _r, lab, _c, _f in rows if lab)
    dupes = {k: v for k, v in seen.items() if v > 1}
    # --- other sheets: NPV on Summary / Cover ---------------------------------------------
    for sh in wb.sheetnames:
        if sh == calcs_sheet:
            continue
        for row in wb[sh].iter_rows():
            for cell in row:
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith("=")
                    and re.search(r"\bNPV\(", cell.value)
                ):
                    findings["4.01-02 NPV function used (never allowed)"].append(
                        (cell.row, f"[{sh}] {wb[sh].cell(row=cell.row, column=2).value}", "")
                    )

    out(f"=== FAST Standard 02c audit — {book}")
    out(f"{len(rows)} formula rows on {calcs_sheet}, {nper} periods\n")
    undeclared_total = 0
    for rule in sorted(findings):
        hits = findings[rule]
        code = rule.split()[0].split("/")[0]
        live = [h for h in hits if (code, re.sub(r"^\[.*?\] ", "", h[1]).strip()) not in declared]
        undeclared_total += len(live)
        out(f"{rule}: {len(hits)} row(s), {len(hits) - len(live)} declared")
        for r, lab, note in hits[:6]:
            key = (code, re.sub(r"^\[.*?\] ", "", lab).strip())
            tag = "  [DECLARED]" if key in declared else ""
            out(f"    row {r:4} {lab[:52]:52s} {note}{tag}")
        if len(hits) > 6:
            out(f"    ... and {len(hits) - 6} more")
        out("")
    out(f"UNDECLARED BREACHES: {undeclared_total}")
    out(f"3.05-03 duplicate labels: {len(dupes)} label(s) used more than once")
    for k, v in list(sorted(dupes.items(), key=lambda x: -x[1]))[:6]:
        out(f"    x{v}  {k[:64]}")
    return FastAudit(dict(findings), undeclared_total, len(rows), nper, dupes)
