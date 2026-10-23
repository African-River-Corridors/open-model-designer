"""Perturbation audit — change each input cell, recalculate, see what actually moves.

Why (2026-09-10): a model's barge-power inputs were changed and the result did not move. A whole
block was Python constants baked into cells. Seventeen inputs were inert and nothing complained.
`verify()` cannot catch it — it asks whether the workbook reproduces the Python oracle, and a
baked constant reproduces it perfectly.

A cheap build-time guard catches an input no formula MENTIONS. This is the strong check, run on
demand: it catches an input a formula mentions but whose value cannot move a single number, for
example one that only appears inside an IF gate.

It writes each perturbed value into a COPY; the source workbook is never touched. It is slow
(~15 minutes on a full model). Redirect output to a file; `head` kills it with SIGPIPE.

A sanity gate runs first: a known-live input is bumped and must move enough rows, or the audit
aborts — an oracle that cannot detect a change it should detect proves nothing.
"""

from __future__ import annotations

import collections
import os
import re
import shutil
import tempfile
from collections.abc import Collection, Mapping
from dataclasses import dataclass

import openpyxl

from ._eval import evaluate_cells


@dataclass
class InputsAudit:
    n_inputs: int
    inert: list[tuple[int, str, float]]
    undeclared: list[tuple[int, str, float]]
    moved_rows: dict[int, int]  # input row -> number of output rows that moved

    @property
    def ok(self) -> bool:
        return not self.undeclared


def read_inputs(ws, label_col: int = 1, value_col: int = 2) -> dict[int, tuple[str, float]]:
    """Every numeric input on the sheet: row -> (label, value)."""
    found = {}
    for r in range(1, ws.max_row + 1):
        lab, val = ws.cell(row=r, column=label_col).value, ws.cell(row=r, column=value_col).value
        if isinstance(val, int | float) and not isinstance(val, bool) and isinstance(lab, str):
            found[r] = (lab, float(val))
    return found


def bump(val: float) -> float:
    """A flag is tested by flipping it; everything else by +10%."""
    if val in (0.0, 1.0):
        return 1.0 - val
    return val * 1.10 if abs(val) > 1e-12 else 0.01


def audit_inputs(
    src: str,
    *,
    inputs_sheet: str,
    watched_sheets: tuple[str, ...],
    gate_label: str,
    declared: Collection[str] | Mapping[str, str] = (),
    label_col: int = 1,
    value_col: int = 2,
    gate_min_rows: int = 20,
    rel_tol: float = 1e-9,
    shard: tuple[int, int] | None = None,
    out=print,
) -> InputsAudit:
    """Bump every input in `inputs_sheet` and count the cells that move on `watched_sheets`.

    `shard=(index, count)` bumps only every `count`-th input from `index` (the sanity gate runs in every shard), so
    `count` runs together cover every input once; CI uses it to run the audit in parallel jobs.

    `gate_label` is the start of a label of a known-live input (the sanity gate). `declared` is
    the set of label prefixes that may legitimately move nothing; any other inert input is a
    defect (`InputsAudit.undeclared`).
    """
    wbk = openpyxl.load_workbook(src)
    if inputs_sheet not in wbk.sheetnames:
        raise AssertionError(
            f"no {inputs_sheet!r} tab in {src} — the sheet names have moved; read them from the "
            f"model's own map, never retype them"
        )
    inputs = read_inputs(wbk[inputs_sheet], label_col, value_col)
    all_inputs = inputs
    if shard is not None:
        inputs = {r: v for i, (r, v) in enumerate(inputs.items()) if i % shard[1] == shard[0]}
    cell_re = re.compile(
        "^(" + "|".join(re.escape(s.upper()) for s in watched_sheets) + r")'!([A-Z]+)(\d+)$"
    )
    tmp = tempfile.mkdtemp(prefix="audit-inputs-")

    def scalars(path: str) -> dict[str, float]:
        return {
            k: v for k, v in evaluate_cells(path, watched_sheets).items() if isinstance(v, float)
        }

    def run(row: int, value: float) -> dict[str, float]:
        path = os.path.join(tmp, f"p{row}.xlsx")
        shutil.copy(src, path)
        w = openpyxl.load_workbook(path)
        w[inputs_sheet].cell(row=row, column=value_col, value=value)
        w.save(path)
        try:
            return scalars(path)
        finally:
            os.remove(path)

    base = scalars(src)

    def moved(got: dict[str, float]) -> collections.Counter:
        mv: collections.Counter = collections.Counter()
        for k, b in base.items():
            g = got.get(k)
            if g is None:
                continue
            if abs(g - b) / max(abs(b), abs(g), 1e-6) > rel_tol:
                m = cell_re.match(k)
                if m:
                    mv[(m.group(1), int(m.group(3)))] += 1
        return mv

    try:
        out(f"baseline: {len(base)} scalar cells on {', '.join(watched_sheets)}")
        gate_row = next(r for r, (lbl, _v) in all_inputs.items() if lbl.startswith(gate_label))
        gate = moved(run(gate_row, all_inputs[gate_row][1] * 1.10))
        out(f"sanity gate — {gate_label!r} +10% moves {len(gate)} rows")
        if not len(gate) > gate_min_rows:
            raise AssertionError(
                "ORACLE BROKEN — a known-live input moved too little. Trust no result below."
            )

        results = {}
        for r, (lab, val) in inputs.items():
            n = len(moved(run(r, bump(val))))
            results[r] = (lab, val, n)
            out(f"  B{r:<3d} {lab[:62]:62s} {n:4d} rows")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    inert = [(r, lb, v) for r, (lb, v, n) in results.items() if n == 0]
    undeclared = [(r, lb, v) for r, lb, v in inert if not any(lb.startswith(w) for w in declared)]
    out(f"\n{len(inputs)} inputs · {len(inert)} inert · {len(undeclared)} of those undeclared")
    for r, lb, v in inert:
        tag = "declared" if (r, lb, v) not in undeclared else "*** UNDECLARED — DEFECT ***"
        out(f"  B{r:<3d} {lb[:62]:62s} {tag}")
    if not undeclared:
        out(
            "AUDIT OK — every input either moves a number or is a declared non-computational input."
        )
    return InputsAudit(len(inputs), inert, undeclared, {r: n for r, (_l, _v, n) in results.items()})
