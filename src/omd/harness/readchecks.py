"""Evaluate a built workbook and read its check block back, the way a reviewer would.

`verify()` proves Calcs reproduces the Python oracle and stops there. This computes the
workbook's own formulas and reads a presentation sheet's check block. It returns non-zero on any
failure: a harness that prints a failure and returns 0 is a sentence, not a gate.

Two check-block conventions exist, so there are two readers:

* `read_status_checks` — a "Validity checks" block with a status word per row (column C) and a
  severity (column D). An `error` row whose status is not in `PASSING` fails.
* `read_zero_checks` — a "Check" block of numbers per period. A row labelled with "≥ 0" must be
  ≥ 0 in every period; every other row must be 0. A value that cannot be read is a failure.

`KNOWN_OPEN` is empty by default. Pass a label only with a reason and a date, and drop it when
the check is fixed.
"""

from __future__ import annotations

import os
import warnings
from collections.abc import Collection

warnings.filterwarnings("ignore")

import formulas  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter as CL  # noqa: E402

PASSING = frozenset({"OK", "APPLIED", "NONE", "OFF", "NOT TRIGGERED", "NOT BINDING"})
TOL = 0.01
FIRST_COL = 3  # column C holds the first period


def read_status_checks(
    path: str,
    *,
    heading: str = "Validity checks",
    passing: Collection[str] = PASSING,
    known_open: Collection[str] = (),
    min_rows: int = 11,
    out=print,
) -> int:
    """Read a status-word check block. Returns 0 if no unexpected error-severity check fails."""
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    wb = openpyxl.load_workbook(path)
    ws = next(
        (
            wb[n]
            for n in wb.sheetnames
            if any(
                isinstance(r[0].value, str) and r[0].value.startswith(heading)
                for r in wb[n].iter_rows(min_col=2, max_col=2)
            )
        ),
        None,
    )
    if not (ws is not None):
        raise AssertionError(f"no sheet carries a {heading!r} block")
    start = next(
        r[0].row
        for r in ws.iter_rows(min_col=2, max_col=2)
        if isinstance(r[0].value, str) and r[0].value.startswith(heading)
    )
    base = os.path.basename(path)

    def val(col, r):
        try:
            return sol[f"'[{base}]{ws.title.upper()}'!{col}{r}"].value[0, 0]
        except Exception:  # noqa: BLE001
            return "?"

    rows, bad, unexpected = 0, [], []
    for r in range(start + 2, start + 120):
        lab = ws.cell(row=r, column=2).value
        if not (isinstance(lab, str) and lab.strip()):
            break
        rows += 1
        v, sev = str(val("C", r)), ws.cell(row=r, column=4).value
        out(f"  {v:14s} {lab[:70]}")
        if v not in passing and sev == "error":
            (bad if lab in known_open else unexpected).append((lab, v))

    if not (rows >= min_rows):
        raise AssertionError(f"the check block scanned {rows} rows — it found almost nothing")
    out(
        f"\n{rows} checks read. known-open failing: {len(bad)}. UNEXPECTED failing: {len(unexpected)}"
    )
    for lab, v in unexpected:
        out(f"   {v}  {lab}")
    return 1 if unexpected else 0


def read_zero_checks(
    path: str,
    sheet: str = "Annual Summary",
    *,
    heading: str = "Check",
    known_open: Collection[str] = (),
    min_rows: int = 10,
    tol: float = TOL,
    out=print,
) -> int:
    """Read a numeric check block. Returns 0 if every row holds in every period."""
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    ws = openpyxl.load_workbook(path)[sheet]
    base = os.path.basename(path)

    start = next((c.row for c in ws["B"] if c.value == heading), None)
    if not (start is not None):
        raise AssertionError(f"{sheet} carries no {heading!r} block")
    ncols = (
        next(
            j
            for j in range(FIRST_COL, ws.max_column + 2)
            if ws.cell(row=start + 1, column=j).value is None
        )
        - FIRST_COL
    )

    def val(r: int, j: int):
        try:
            return sol[f"'[{base}]{sheet.upper()}'!{CL(j)}{r}"].value[0, 0]
        except Exception:  # noqa: BLE001
            return None

    rows, bad, unexpected = 0, [], []
    r = start + 1
    while (
        isinstance(ws.cell(row=r, column=2).value, str) and ws.cell(row=r, column=2).value.strip()
    ):
        lab = ws.cell(row=r, column=2).value
        vals = [val(r, j) for j in range(FIRST_COL, FIRST_COL + ncols)]
        nonneg = "≥ 0" in lab
        fails = [
            v
            for v in vals
            if not isinstance(v, int | float) or (v < -tol if nonneg else abs(v) > tol)
        ]
        worst = max((abs(v) for v in vals if isinstance(v, int | float)), default=float("nan"))
        out(f"  {'OK' if not fails else 'FAIL':5s} {lab[:72]:72s} max|v| {worst:,.2f}")
        if fails:
            (bad if lab in known_open else unexpected).append(lab)
        rows += 1
        r += 1

    if not (rows >= min_rows):
        raise AssertionError(f"the check block scanned {rows} rows — it found almost nothing")
    if not (ncols >= 2):
        raise AssertionError(f"the check block spans {ncols} period columns")
    out(
        f"\n{rows} checks × {ncols} periods read. known-open failing: {len(bad)}. "
        f"UNEXPECTED failing: {len(unexpected)}"
    )
    for lab in unexpected:
        out(f"   FAIL  {lab}")
    return 1 if unexpected else 0


def main(argv: list[str] | None = None) -> int:
    """`python -m omd.harness.readchecks status|zero WORKBOOK [--sheet S] [--min-rows N]`."""
    import argparse

    ap = argparse.ArgumentParser(prog="omd.harness.readchecks")
    ap.add_argument("kind", choices=("status", "zero"))
    ap.add_argument("path")
    ap.add_argument("--sheet", default="Annual Summary", help="zero: the sheet holding the block")
    ap.add_argument("--heading", default=None)
    ap.add_argument("--min-rows", type=int, default=None)
    a = ap.parse_args(argv)
    if a.kind == "status":
        kw = {"heading": a.heading} if a.heading else {}
        if a.min_rows is not None:
            kw["min_rows"] = a.min_rows
        return read_status_checks(a.path, **kw)
    kw = {"heading": a.heading} if a.heading else {}
    if a.min_rows is not None:
        kw["min_rows"] = a.min_rows
    return read_zero_checks(a.path, a.sheet, **kw)


if __name__ == "__main__":
    import sys

    sys.exit(main())
