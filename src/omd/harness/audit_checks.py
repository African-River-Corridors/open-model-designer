"""Corruption audit — inject a wrong number and check that some validity check notices.

The lesson (2026-09-10): a balance check was blind to a whole waterfall, because cash and retained
earnings both subtract the same line and it cancels on both sides. A check that cannot fail is
not evidence.

The sweep scales one Calcs cell by `factor` (+25%) — in the period where that row is LARGEST,
because scaling a zero cell is a no-op that looks like an uncaught error — and reports which
checks fire. Every target must be caught by at least one check. A check that already fires on the
clean workbook is not counted as catching anything.

A model supplies its targets and a `fired` function that says which checks fire on an evaluated
workbook. Two builders cover the two conventions:

* `status_checks` — checks are rows on a summary sheet that read a status word such as "FAIL".
* `value_checks` — checks are numeric rows on Calcs with a per-check predicate.

Run it after adding or changing any Calcs row. A row with no check watching it is a row where a
wrong number reaches a board paper silently.
"""

from __future__ import annotations

import os
import shutil
import tempfile
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils import get_column_letter as CL

from ._eval import cell, evaluate_cells

#: `fired(values)` -> names of the checks that fire on an evaluated workbook.
Fired = Callable[[dict], list[str]]


@dataclass
class AuditResult:
    caught: dict[str, list[str]] = field(default_factory=dict)  # target -> checks that fired
    uncaught: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)  # zero in every period
    missing: list[str] = field(default_factory=list)  # row not found
    baseline_fired: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        """True only if something was tested and everything tested was caught. An audit that
        skipped every target, found no target, or ran on a book whose checks were already red
        proved nothing."""
        return (
            bool(self.caught)
            and not self.uncaught
            and not self.missing
            and not self.skipped
            and not self.baseline_fired
        )


def label_rows(ws) -> dict[str, int]:
    """Stripped column-A label -> row. A duplicate label keeps its FIRST row."""
    rows: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip():
            rows.setdefault(v.strip(), r)
    return rows


def find_row(rows: Mapping[str, int], label: str) -> int | None:
    """Exact (stripped) label first, else the first label that starts with it."""
    want = label.strip()
    if want in rows:
        return rows[want]
    return next((r for lbl, r in rows.items() if lbl.startswith(want)), None)


def status_checks(
    src: str,
    *,
    sheet: str = "Summary",
    label_col: int = 2,
    status_col: str = "C",
    fail: str = "FAIL",
    exclude: Sequence[str] = ("check",),
) -> Fired:
    """Checks that read a status word. Every labelled row on `sheet` with a value in the status
    column is a check, except the labels in `exclude`."""
    ws = openpyxl.load_workbook(src)[sheet]
    checks = {
        ws.cell(row=r, column=label_col).value: r
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=label_col).value, str)
        and ws.cell(row=r, column=column_index_from_string(status_col)).value
    }
    watch = {k: r for k, r in checks.items() if k not in exclude}

    def fired(values: dict) -> list[str]:
        return [k for k, r in watch.items() if str(cell(values, sheet, f"{status_col}{r}")) == fail]

    return fired


def value_checks(
    src: str,
    spec: Mapping[str, tuple[str, Callable[[float], bool]]],
    *,
    calcs_sheet: str = "Calcs",
    first_col: int = 3,
    n_periods: int | None = None,
) -> Fired:
    """Checks that are numeric rows on Calcs. `spec` maps a check name to (label prefix,
    predicate); the predicate is True when the check fires on a value."""
    ws = openpyxl.load_workbook(src)[calcs_sheet]
    n = n_periods if n_periods is not None else ws.max_column - (first_col - 1)
    rows = label_rows(ws)
    located = {}
    for name, (prefix, pred) in spec.items():
        r = next((rr for lbl, rr in rows.items() if lbl.startswith(prefix.strip())), None)
        if not r:
            raise AssertionError(f"check row not found: {prefix}")
        located[name] = (r, pred)

    def fired(values: dict) -> list[str]:
        out = []
        for name, (r, pred) in located.items():
            for i in range(n):
                v = cell(values, calcs_sheet, f"{CL(first_col + i)}{r}")
                if isinstance(v, float) and pred(v):
                    out.append(name)
                    break
        return out

    return fired


def corruption_audit(
    src: str,
    targets: Sequence[str],
    fired: Fired,
    *,
    calcs_sheet: str = "Calcs",
    first_col: int = 3,
    n_periods: int | None = None,
    factor: float = 1.25,
    min_magnitude: float = 1.0,
    out=print,
) -> AuditResult:
    """Corrupt each target row in its largest period and record which checks fire."""
    ws = openpyxl.load_workbook(src)[calcs_sheet]
    n = n_periods if n_periods is not None else ws.max_column - (first_col - 1)
    rows = label_rows(ws)
    base = evaluate_cells(src)
    res = AuditResult(baseline_fired=fired(base))
    out(f"baseline checks already firing: {res.baseline_fired or 'none'}\n")

    tmp = tempfile.mkdtemp(prefix="audit-checks-")
    try:
        for t in targets:
            rr = find_row(rows, t)
            if rr is None:
                out(f"  ?? row not found: {t}")
                res.missing.append(t)
                continue
            mags = [
                abs(v)
                if isinstance(v := cell(base, calcs_sheet, f"{CL(first_col + i)}{rr}"), float)
                else 0.0
                for i in range(n)
            ]
            i = max(range(n), key=lambda j: mags[j])
            if mags[i] < min_magnitude:
                out(f"  {t.strip()[:58]:58s} SKIPPED — zero in every period")
                res.skipped.append(t)
                continue
            path = os.path.join(tmp, "corrupt.xlsx")
            shutil.copy(src, path)
            w = openpyxl.load_workbook(path)
            ref = f"{CL(first_col + i)}{rr}"
            old = w[calcs_sheet][ref].value
            w[calcs_sheet][ref] = "=(" + str(old)[1:] + f")*{factor:g}"
            w.save(path)
            try:
                f = [x for x in fired(evaluate_cells(path)) if x not in res.baseline_fired]
            finally:
                os.remove(path)
            status = "CAUGHT by " + ", ".join(f"'{x}'" for x in f) if f else "*** NOT CAUGHT ***"
            out(f"  {t.strip()[:58]:58s} x{factor:g} -> {status}")
            if f:
                res.caught[t] = f
            else:
                res.uncaught.append(t)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    done = len(res.caught) + len(res.uncaught)
    out(
        f"\n{len(res.caught)} of {done} injected errors caught ({len(res.skipped)} skipped as all-zero)"
    )
    if res.uncaught:
        out("UNCAUGHT:")
        for b in res.uncaught:
            out(f"  - {b.strip()}")
    if res.missing:
        out(f"{len(res.missing)} target row(s) not found")
    return res
