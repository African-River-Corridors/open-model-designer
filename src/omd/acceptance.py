"""Acceptance CI for library modules.

A module is accepted only if it ships all six (the sixth is the FAST audit):

1. ports declared, with units;
2. oracle and formula halves agree on a worked example, and the formulas follow the input cells
   (change an input cell, re-evaluate, compare with the oracle: a baked-in number fails);
3. a golden case from an independent source, not the module's own oracle;
4. a clean corruption audit: every row the module writes is caught by one of its checks;
5. a scope note: what it covers and what it does not;
6. the FAST audit of its worked-example book has no undeclared finding on a calculation row.

Independence of a golden case is declared (a kind and a named source) and checked only for not
being the module's own oracle or worked example; a reviewer must read the source.

`check_module` returns every failure, each tagged with the criterion it breaks. `assert_accepted`
raises `AcceptanceError` listing them. Run it from a library's CI:

    python -m omd.acceptance package.module:MODULE
    python -m omd.acceptance --all package.modules     # every module in the package (T0)
    python -m omd.acceptance --all package.modules --shard 1/4   # shard 1 (of 0..3) of 4, for CI
"""

from __future__ import annotations

import importlib
import multiprocessing
import os
import sys
import tempfile
from dataclasses import dataclass

import openpyxl

from .harness import fastblock
from .harness._eval import cell, evaluate_cells
from .lines import STANDARD_LINES, dimension_of, valid_unit
from .module import INDEPENDENT_KINDS, BuildContext, Module
from .sheet.verify import evaluate, verify_values
from .sheet.workbook import Workbook

CRITERIA = {
    "1": "ports declared, with units",
    "2": "oracle and formula halves agree on a worked example",
    "3": "golden case from an independent source",
    "4": "clean corruption audit",
    "5": "scope note",
    "6": "FAST audit of the worked-example book",
}
N_PERIODS = 4
#: FAST rules a check row is exempt from: a check re-derives a row from the ports in one formula (3.03-01), and in
#: period 1 it reads the opening input, so the first column's shape differs (3.02-01). Every other rule applies to
#: check rows too.
CHECK_ROW_EXEMPT = ("3.03-01", "3.02-01")
CHECK_PREFIX = "Check -"
TOL = 1e-6


@dataclass(frozen=True)
class Failure:
    criterion: str
    message: str

    def __str__(self) -> str:
        return f"criterion {self.criterion} ({CRITERIA[self.criterion]}): {self.message}"


class AcceptanceError(AssertionError):
    def __init__(self, failures: list[Failure]):
        self.failures = failures
        super().__init__("\n".join(str(f) for f in failures))

    @property
    def criteria(self) -> set[str]:
        return {f.criterion for f in self.failures}


# --- 1. ports ------------------------------------------------------------------------------------
def check_ports(m: Module) -> list[Failure]:
    out: list[Failure] = []
    f = lambda msg: out.append(Failure("1", msg))  # noqa: E731
    if not m.inputs:
        f("no input ports declared")
    if not m.outputs:
        f("no output ports declared")
    names = [p.name for p in m.ports]
    for dup in sorted({n for n in names if names.count(n) > 1}):
        f(f"port name {dup!r} is used twice")
    for p in m.ports:
        if p.kind not in ("series", "scalar"):
            f(f"port {p.name!r}: kind must be 'series' or 'scalar', got {p.kind!r}")
        if not valid_unit(p.unit):
            f(f"port {p.name!r}: missing or malformed unit {p.unit!r}")
            continue
        line = p.standard_line
        if p.line is not None and p.line not in STANDARD_LINES:
            f(f"port {p.name!r}: {p.line!r} is not a standard line name")
        elif line is not None and dimension_of(p.unit) != STANDARD_LINES[line]:
            f(
                f"port {p.name!r}: unit {p.unit!r} is {dimension_of(p.unit)}, "
                f"but standard line {line!r} is {STANDARD_LINES[line]}"
            )
    return out


# --- building the module into a workbook --------------------------------------------------------
def _host(m: Module, inputs: dict, path: str):
    """A minimal host model: one Inputs sheet for scalar ports, one Feed sheet for series ports."""
    fastblock.reset()
    n = N_PERIODS
    wb = Workbook(f"{m.name} acceptance", [f"P{i + 1}" for i in range(n)], opening=True)
    scalars = {p.name: (p.name, float(inputs[p.name])) for p in m.inputs if p.kind == "scalar"}
    refs = wb.inputs("Scalar ports", scalars) if scalars else {}
    sources = dict(refs)
    series = [p for p in m.inputs if p.kind == "series"]
    if series:
        feed = wb.add_sheet("Feed")
        for r, p in enumerate(series, start=2):
            feed.cell(row=r, column=1, value=p.name)
            for i, v in enumerate(inputs[p.name]):
                feed.cell(row=r, column=3 + i, value=float(v))
            sources[p.name] = f"Feed!{{col}}{r}"
    calcs = wb.calcs()
    ctx = BuildContext(m, calcs, n, sources, inputs)
    first = calcs.next_row
    m.build(ctx)
    spec = m.checks(ctx)
    last = calcs.next_row  # the span covers the module's check rows too
    wb.save(path)
    return calcs, ctx, spec, (first, last)


def _expected(m: Module, inputs: dict) -> dict[str, list[float]]:
    return {k: [float(x) for x in v] for k, v in m.oracle(inputs).items()}


def _compare(m: Module, inputs: dict, want: dict, criterion: str, label: str) -> list[Failure]:
    """Build the formula half on `inputs`; its outputs must equal `want`."""
    out: list[Failure] = []
    declared = {p.name for p in m.outputs}
    if set(want) != declared:
        return [
            Failure(
                criterion, f"{label}: expected outputs {sorted(want)} != ports {sorted(declared)}"
            )
        ]
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.xlsx")
        try:
            calcs, ctx, _spec, _span = _host(m, inputs, path)
        except Exception as e:  # noqa: BLE001 — a module that cannot build fails here
            return [Failure(criterion, f"{label}: the formula half did not build: {e}")]
        rows = calcs.rows()
        for p in m.outputs:
            if ctx.key(p.name) not in rows:
                out.append(Failure(criterion, f"{label}: no row registered for output {p.name!r}"))
        if out:
            return out
        values = evaluate(path)  # one evaluation checks every output
        for p in m.outputs:
            try:
                verify_values(values, calcs, {ctx.key(p.name): want[p.name]}, tol=TOL)
            except AssertionError as e:
                out.append(Failure(criterion, f"{label}: output {p.name!r}: {e}"))
    return out


# --- 2. oracle vs formulas on the worked example --------------------------------------------------
def check_worked_example(m: Module) -> list[Failure]:
    try:
        want = _expected(m, m.example)
    except Exception as e:  # noqa: BLE001
        return [Failure("2", f"the oracle failed on the worked example: {e}")]
    fails = _compare(m, m.example, want, "2", "worked example")
    return fails or check_inputs_drive_outputs(m)


def _bump(v):
    """A different value for an input: +10%, or 1 if it was 0."""
    if isinstance(v, list | tuple):
        return [_bump(x) for x in v]
    return float(v) * 1.1 if float(v) != 0 else 1.0


def check_inputs_drive_outputs(m: Module) -> list[Failure]:
    """Change each input CELL in the built workbook and re-evaluate: the outputs must equal the
    oracle on the changed inputs. A formula half that carries a number it took from the inputs at
    build time agrees with the oracle on every case it was built for, and fails here."""
    out: list[Failure] = []
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.xlsx")
        try:
            calcs, ctx, _spec, _span = _host(m, m.example, path)
        except Exception as e:  # noqa: BLE001
            return [Failure("2", f"the formula half did not build: {e}")]
        for p in m.inputs:
            moved = {**m.example, p.name: _bump(m.example[p.name])}
            try:
                want = _expected(m, moved)
            except Exception as e:  # noqa: BLE001
                out.append(Failure("2", f"the oracle failed on changed input {p.name!r}: {e}"))
                continue
            w = openpyxl.load_workbook(path)
            src = ctx.source(p.name)
            sheet, ref = src.split("!", 1)
            if p.kind == "scalar":
                w[sheet][ref.replace("$", "")] = moved[p.name]
            else:
                row = int(ref.removeprefix("{col}"))
                for i, v in enumerate(moved[p.name]):
                    w[sheet].cell(row=row, column=3 + i, value=float(v))
            changed = os.path.join(d, f"moved-{p.name}.xlsx")
            w.save(changed)
            values = evaluate(changed)  # one evaluation checks every output
            for o in m.outputs:
                try:
                    verify_values(values, calcs, {ctx.key(o.name): want[o.name]}, tol=TOL)
                except AssertionError as e:
                    out.append(Failure(
                        "2", f"output {o.name!r} does not follow input {p.name!r} (a value baked "
                        f"into the formulas?): {e}"))  # fmt: skip
    return out


# --- 3. golden case ----------------------------------------------------------------------------------
def check_golden(m: Module) -> list[Failure]:
    g = m.golden
    out: list[Failure] = []
    if g is None or not (g.source or "").strip():
        return [Failure("3", "no golden case, or it names no source")]
    if g.kind not in INDEPENDENT_KINDS:
        return [
            Failure(
                "3",
                f"golden kind {g.kind!r} is not independent; use one of {sorted(INDEPENDENT_KINDS)}",
            )
        ]
    if g.inputs == m.example:
        out.append(
            Failure("3", "the golden case is the worked example; it must be a different case")
        )
    if set(g.outputs) != {p.name for p in m.outputs}:
        return [
            *out,
            Failure("3", "the golden case must give an expected value for every output port"),
        ]
    try:
        got = _expected(m, g.inputs)
    except Exception as e:  # noqa: BLE001
        return [*out, Failure("3", f"the oracle failed on the golden case: {e}")]
    for name, want in g.outputs.items():
        if any(abs(a - b) > TOL for a, b in zip(got[name], want, strict=True)):
            out.append(
                Failure(
                    "3",
                    f"the oracle disagrees with the golden case on {name!r}: {got[name]} vs {want}",
                )
            )
    return out or _compare(
        m, g.inputs, {k: list(v) for k, v in g.outputs.items()}, "3", "golden case"
    )


# --- 4. corruption audit ----------------------------------------------------------------------------
OPEN_COL = "B"  # the host's Open column; the periods follow
PERIOD_COLS = ("C", "D", "E", "F")


def _fired(values: dict, check_rows: dict[str, int]) -> list[str]:
    """Checks that read non-zero in any period."""
    out = []
    for name, r in check_rows.items():
        for col in PERIOD_COLS:
            v = cell(values, "Calcs", f"{col}{r}")
            if isinstance(v, float) and abs(v) > 1e-6:
                out.append(name)
                break
    return out


_C4: dict = {}  # what a corruption worker needs, set before the workers fork


def _corrupt_one(rc: tuple[int, str]) -> Failure | None:
    """Corrupt one cell (row, column) of the built module by 25% and say whether some check caught it."""
    r, col = rc
    ws, path, base = _C4["ws"], _C4["path"], _C4["base"]
    old = ws[f"{col}{r}"].value
    v = cell(base, "Calcs", f"{col}{r}")
    if v is None and isinstance(old, int | float):
        v = float(old)
    delta = 0.25 * max(abs(v) if isinstance(v, float) else 1.0, 1.0)
    w = openpyxl.load_workbook(path)
    if isinstance(old, str):
        w["Calcs"][f"{col}{r}"] = f"=({old[1:]})+{delta!r}"
    else:
        w["Calcs"][f"{col}{r}"] = old + delta
    bad = os.path.join(_C4["dir"], f"corrupt-{os.getpid()}.xlsx")
    w.save(bad)
    fired = [
        x
        for x in _fired(evaluate_cells(bad, ("Calcs",)), _C4["check_rows"])
        if x not in _C4["clean_red"]
    ]
    if fired:
        return None
    where = "opening value" if col == OPEN_COL else f"period {PERIOD_COLS.index(col) + 1}"
    label = _C4["labels"][r]
    return Failure(
        "4", f"row {label or r!r} {where}: written, but no check catches a wrong value in it"
    )


def check_corruption(m: Module) -> list[Failure]:
    """Every cell the module writes (each period, and the Open column) is corrupted in turn, by ROW NUMBER, and
    some check of the module must fire. Rows are not found by label: a duplicate or blank label
    cannot hide a row."""
    out: list[Failure] = []
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.xlsx")
        try:
            calcs, _ctx, spec, (first, last) = _host(m, m.example, path)
        except Exception as e:  # noqa: BLE001
            return [Failure("4", f"the formula half did not build: {e}")]
        if not spec:
            return [Failure("4", "the module declares no checks")]
        ws = calcs.ws
        labels = {r: str(ws.cell(row=r, column=1).value or "").strip() for r in range(first, last)}
        check_rows: dict[str, int] = {}
        for name, prefix in spec.items():
            hit = next((r for r, lab in labels.items() if lab.startswith(prefix.strip())), None)
            if hit is None:
                return [Failure("4", f"check {name!r}: no row labelled {prefix!r}")]
            check_rows[name] = hit

        def written_cols(r):
            """The cells of a row the module wrote: its period formulas and its opening cell."""
            cols = [
                c for c in PERIOD_COLS
                if isinstance(ws[f"{c}{r}"].value, str) and ws[f"{c}{r}"].value.startswith("=")
            ]  # fmt: skip
            if ws[f"{OPEN_COL}{r}"].value is not None:
                cols.insert(0, OPEN_COL)
            return cols

        targets = [
            r for r in range(first, last) if written_cols(r) and r not in check_rows.values()
        ]
        if not targets:
            return [Failure("4", "the module writes no rows to audit")]
        base = evaluate_cells(path, ("Calcs",))
        clean_red = _fired(base, check_rows)
        if clean_red:
            out.append(Failure("4", f"checks fire on the clean worked example: {clean_red}"))
        for r in targets:
            if not labels[r]:
                out.append(Failure("4", f"row {r} is written but has no label"))
        cells = [(r, col) for r in targets for col in written_cols(r)]
        global _C4
        _C4 = {"ws": ws, "path": path, "base": base, "check_rows": check_rows, "clean_red": clean_red,
               "dir": d, "labels": labels}  # fmt: skip
        workers = int(os.environ.get("OMD_AUDIT_WORKERS", "1"))
        if workers > 1 and "fork" in multiprocessing.get_all_start_methods():
            from concurrent.futures import ProcessPoolExecutor

            with ProcessPoolExecutor(
                workers, mp_context=multiprocessing.get_context("fork")
            ) as pool:
                results = list(pool.map(_corrupt_one, cells, chunksize=4))
        else:
            results = [_corrupt_one(rc) for rc in cells]
        out.extend(f for f in results if f is not None)
    return out


# --- 5. scope note --------------------------------------------------------------------------------------
def check_scope(m: Module) -> list[Failure]:
    s = m.scope
    out = []
    if not isinstance(s, dict):
        return [Failure("5", "scope must be a mapping with 'in' and 'out'")]
    for key in ("in", "out"):
        items = s.get(key)
        if not (
            isinstance(items, list)
            and items
            and all(isinstance(i, str) and i.strip() for i in items)
        ):
            out.append(
                Failure(
                    "5",
                    f"scope needs a non-empty list under {key!r} (what it covers / does not cover)",
                )
            )
    return out


# --- 6. FAST audit ------------------------------------------------------------------------------------
def check_fast(m: Module) -> list[Failure]:
    """The FAST audit (`omd.harness.audit_fast`) on the worked-example book. A finding on a calculation row is a
    failure unless the module declares it in `fast_declared` ({"<rule code> <label prefix>": "reason"}). Check rows are
    exempt from `CHECK_ROW_EXEMPT` only."""
    from .harness.audit_fast import audit_fast

    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.xlsx")
        try:
            _host(m, m.example, path)
        except Exception as e:  # noqa: BLE001
            return [Failure("6", f"the formula half did not build: {e}")]
        audit = audit_fast(path, first_col=3, out=lambda *_: None)
    out: list[Failure] = []
    declared = []
    for key, reason in (m.fast_declared or {}).items():
        code, _, prefix = key.partition(" ")
        if not (code and prefix.strip() and str(reason).strip()):
            out.append(
                Failure(
                    "6", f"fast_declared {key!r}: needs '<rule code> <label prefix>' and a reason"
                )
            )
        declared.append((code, prefix.strip(), key))
    used: set[str] = set()
    for rule, hits in audit.findings.items():
        code = rule.split()[0].split("/")[0]
        for _row, label, _note in hits:
            label = label.strip()
            if label.startswith(CHECK_PREFIX) and code in CHECK_ROW_EXEMPT:
                continue
            match = next(
                (k for c, prefix, k in declared if c == code and label.startswith(prefix)), None
            )
            if match is not None:
                used.add(match)
                continue
            out.append(Failure("6", f"row {label!r}: FAST {rule}"))
    for _c, _p, key in declared:
        if key not in used:
            out.append(
                Failure(
                    "6", f"fast_declared {key!r} matches no finding: remove the stale declaration"
                )
            )
    return out


CHECKS = {
    "1": check_ports,
    "2": check_worked_example,
    "3": check_golden,
    "4": check_corruption,
    "5": check_scope,
    "6": check_fast,
}


def check_module(m: Module, criteria: set[str] | None = None) -> list[Failure]:
    """Every failure of every criterion asked for (default all). An empty list is acceptance."""
    out: list[Failure] = []
    for c, fn in CHECKS.items():
        if criteria is None or c in criteria:
            out.extend(fn(m))
    return out


def assert_accepted(m: Module, criteria: set[str] | None = None) -> None:
    failures = check_module(m, criteria)
    if failures:
        raise AcceptanceError(failures)


def _report(m: Module) -> int:
    failures = check_module(m)
    for f in failures:
        print(f"FAIL {m.name}: {f}")
    if failures:
        print(f"{m.name} {m.version}: NOT ACCEPTED ({len(failures)} failure(s))")
        return 1
    print(
        f"{m.name} {m.version}: accepted (criteria 1-6); {len(m.fast_declared or {})} FAST declaration(s)"
    )
    return 0


def _check_one(package: str, name: str) -> tuple[str, str, list[str], int, float]:
    """In a worker process: discover the package again and check one module by name."""
    import time

    from .registry import discover

    m = discover(package)[name]
    if (
        SHARD_WEIGHT.get(name, 10) >= 100
    ):  # a heavy module's corruption audit uses every core (it runs nearly alone)
        os.environ["OMD_AUDIT_WORKERS"] = str(os.cpu_count() or 1)
    t0 = time.monotonic()
    failures = [str(f) for f in check_module(m)]
    return name, m.version, failures, len(m.fast_declared or {}), time.monotonic() - t0


def _check_all(package: str, names: list[str]) -> list[str]:
    """Check every module, one worker process each; print each report whole, in module order.
    Returns the names that were not accepted."""
    from concurrent.futures import ProcessPoolExecutor

    bad: list[str] = []
    total_declared = 0
    with ProcessPoolExecutor() as pool:
        for name, version, failures, ndecl, secs in pool.map(
            _check_one, [package] * len(names), names
        ):
            total_declared += ndecl
            for f in failures:
                print(f"FAIL {name}: {f}")
            if failures:
                print(f"{name} {version}: NOT ACCEPTED ({len(failures)} failure(s)); {secs:.0f} s")
                bad.append(name)
            else:
                print(
                    f"{name} {version}: accepted (criteria 1-6); {ndecl} FAST declaration(s); {secs:.0f} s"
                )
    print(f"{total_declared} FAST declaration(s) in total")
    return bad


#: seconds a module's acceptance takes on a laptop, for dealing modules into shards; a module not listed counts as 10.
#: Only the balance of the shards depends on it, never which modules are checked.
SHARD_WEIGHT = {"three_statements": 160, "tax_losses": 50, "cash_waterfall": 40, "scenarios": 240}


def shard_names(names: list[str], count: int, index: int) -> list[str]:
    """The modules of shard `index` of `count`: heaviest first to the lightest shard; every name in exactly one shard."""
    loads = [0] * count
    mine = []
    for name in sorted(names, key=lambda n: (-SHARD_WEIGHT.get(n, 10), n)):
        i = loads.index(min(loads))
        loads[i] += SHARD_WEIGHT.get(name, 10)
        if i == index:
            mine.append(name)
    return [n for n in names if n in mine]


def main(argv: list[str] | None = None) -> int:
    args = sys.argv[1:] if argv is None else argv
    shard = None
    if len(args) == 4 and args[0] == "--all" and args[2] == "--shard":
        index, count = (int(x) for x in args[3].split("/"))
        shard = (count, index)
        args = args[:2]
    if len(args) == 2 and args[0] == "--all":
        from .registry import DiscoveryError, discover

        try:
            mods = discover(args[1])
        except DiscoveryError as e:
            print(f"FAIL discovery: {e}")
            return 1
        if not mods:
            print(f"FAIL discovery: no modules found in {args[1]}")
            return 1
        names = list(mods) if shard is None else shard_names(list(mods), shard[0], shard[1])
        if not names:
            print(f"FAIL discovery: shard {shard[1]}/{shard[0]} found no modules")
            return 1
        bad = _check_all(args[1], names)
        print(f"{len(names) - len(bad)} of {len(names)} modules accepted")
        return 1 if bad else 0
    if len(args) != 1 or ":" not in args[0]:
        print("usage: python -m omd.acceptance package.module:MODULE | --all package.modules")
        return 2
    mod, attr = args[0].split(":")
    return _report(getattr(importlib.import_module(mod), attr))


if __name__ == "__main__":
    sys.exit(main())
