"""Compose modules into one workbook: the engine of a template.

A template is a list of `Step`s. Each step is a module and a wiring that says where each input port reads:
`"ext:<name>"` is an external input (a scalar on the Inputs sheet, a series on the Feed sheet), and
`"<module>.<port>"` is the output row of an earlier step. Steps are built in the order given, so a step may read only
earlier steps. Each module keeps its own formula half, checks and oracle: the template adds the wiring, one Summary
sheet that gathers every module's check rows into one block (read by `read_zero_checks`), and a chained oracle.

`compose` returns a `Composed` that can be saved, evaluated against the chained oracle, and audited.
"""

from __future__ import annotations

import os
import tempfile
from dataclasses import dataclass, field

import openpyxl

from . import acceptance
from .harness import fastblock
from .harness._eval import evaluate_cells
from .module import BuildContext, Module
from .sheet.workbook import Workbook

N_PERIODS = acceptance.N_PERIODS


@dataclass(frozen=True)
class Step:
    module: Module
    wiring: dict[str, str]  # input port -> "ext:<name>" or "<module name>.<port>"


@dataclass
class Composed:
    wb: Workbook
    calcs: object
    steps: list[Step]
    ctxs: dict[str, BuildContext]
    specs: dict[str, dict[str, str]]
    spans: dict[str, tuple[int, int]]
    expected: dict[str, dict[str, list[float]]]  # module name -> chained oracle outputs
    check_rows: list[tuple[str, str, int]] = field(
        default_factory=list
    )  # (module, label, Calcs row)

    def save(self, path: str) -> None:
        self.wb.save(path)


def chained_oracle(steps: list[Step], external: dict) -> dict[str, dict[str, list[float]]]:
    """Run every module's oracle in order, feeding each its inputs from the external values or earlier outputs."""
    out: dict[str, dict[str, list[float]]] = {}
    for st in steps:
        out[st.module.name] = st.module.oracle(_inputs_of(st, external, out))
    return out


def _inputs_of(st: Step, external: dict, out: dict) -> dict:
    vals = {}
    for p in st.module.inputs:
        src = st.wiring[p.name]
        if src.startswith("ext:"):
            vals[p.name] = external[src[4:]]
        else:
            mod, port = src.split(".", 1)
            v = out[mod][port]
            vals[p.name] = (
                v[0] if p.kind == "scalar" else v
            )  # a scalar port reads the first period of a series row
    return vals


def compose(
    steps: list[Step], external: dict, title: str = "Template", *, link_checks=()
) -> Composed:
    """`link_checks`: (label, formula(comp, col)) pairs, each a check row on Calcs that joins modules the ports cannot
    (for example a scalar that must equal the last value of a series). Each reads 0 when it holds."""
    names = [s.module.name for s in steps]
    assert len(set(names)) == len(names), "a module may appear once in a template"
    fastblock.reset()
    n = N_PERIODS
    expected = chained_oracle(steps, external)
    wb = Workbook(title, [f"P{i + 1}" for i in range(n)], opening=True)

    ext_series: dict[str, str] = {}
    ext_scalar: dict[str, str] = {}
    for st in steps:
        for p in st.module.inputs:
            src = st.wiring[p.name]
            if src.startswith("ext:"):
                (ext_series if p.kind == "series" else ext_scalar)[src[4:]] = p.unit
    refs = (
        wb.inputs("Scalar inputs", {k: (k, float(external[k])) for k in ext_scalar})
        if ext_scalar
        else {}
    )
    sources_ext = dict(refs)
    if ext_series:
        feed = wb.add_sheet("Feed")
        for r, k in enumerate(ext_series, start=2):
            feed.cell(row=r, column=1, value=k)
            for i, v in enumerate(external[k]):
                feed.cell(row=r, column=3 + i, value=float(v))
            sources_ext[k] = f"Feed!{{col}}{r}"
    calcs = wb.calcs()

    ctxs: dict[str, BuildContext] = {}
    specs: dict[str, dict[str, str]] = {}
    spans: dict[str, tuple[int, int]] = {}
    for st in steps:
        m = st.module
        sources = {}
        for p in m.inputs:
            src = st.wiring[p.name]
            if src.startswith("ext:"):
                sources[p.name] = sources_ext[src[4:]]
            else:
                mod, port = src.split(".", 1)
                row = calcs.rows()[f"{mod}.{port}"]
                if p.kind == "scalar":  # the first period of the row
                    sources[p.name] = f"{calcs.ws.title}!${wb.period_cols[0]}${row}"
                else:
                    sources[p.name] = f"{calcs.ws.title}!{{col}}{row}"
        ctx = BuildContext(m, calcs, n, sources, _inputs_of(st, external, expected))
        first = calcs.next_row
        m.build(ctx)
        specs[m.name] = m.checks(ctx)
        spans[m.name] = (first, calcs.next_row)
        ctxs[m.name] = ctx

    comp = Composed(wb, calcs, steps, ctxs, specs, spans, expected)
    comp.ext_cells = {
        **{k: ("Inputs", ref.split("!", 1)[1].replace("$", "")) for k, ref in refs.items()},
        **{
            k: ("Feed", int(ref.rsplit("}", 1)[1]))
            for k, ref in sources_ext.items()
            if k in ext_series
        },
    }
    if link_checks:
        calcs.skip_rows()
        calcs.section("Template checks")
        first = calcs.next_row
        for i, (label, fn) in enumerate(link_checks):
            calcs.row(
                f"tpl_chk_{i}",
                f"  {label} (USD)",
                lambda x, fn=fn: "=" + fn(comp, x.col),
                values=[0.0] * n,
            )
        spans["template"] = (first, calcs.next_row)
        specs["template"] = {
            f"tpl{i}": f"  {label}".strip() for i, (label, _fn) in enumerate(link_checks)
        }
    _summary(comp)
    return comp


def _summary(comp: Composed) -> None:
    """One Summary sheet: a 'Check' block with a row per module check, each reading the module's own check row."""
    ws = comp.calcs.ws
    cols = list(comp.wb.period_cols)
    sheet = comp.wb.add_sheet("Summary")
    sheet.cell(row=2, column=2, value="Check")
    r = 3
    for name in comp.specs:
        first, last = comp.spans[name]
        labels = {
            row: str(ws.cell(row=row, column=1).value or "").strip() for row in range(first, last)
        }
        for _key, prefix in comp.specs[name].items():
            row = next(rw for rw, lab in labels.items() if lab.startswith(prefix.strip()))
            sheet.cell(row=r, column=2, value=f"{name}: {labels[row]}")
            for i, col in enumerate(cols):
                sheet.cell(row=r, column=3 + i, value=f"={ws.title}!{col}{row}")
            comp.check_rows.append((name, labels[row], row))
            r += 1


def evaluate_and_compare(comp: Composed, tol: float = 1e-6) -> list[str]:
    """Build, evaluate the workbook's own formulas, and compare every module output row with the chained oracle."""
    problems: list[str] = []
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        comp.save(path)
        values = evaluate_cells(path, ("Calcs",))
    rows = comp.calcs.rows()
    cols = list(comp.wb.period_cols)
    for name, outs in comp.expected.items():
        for port, want in outs.items():
            row = rows[f"{name}.{port}"]
            for i, col in enumerate(cols):
                got = values.get(f"CALCS'!{col}{row}")
                if not isinstance(got, float) or abs(got - want[i]) > tol * max(1.0, abs(want[i])):
                    problems.append(f"{name}.{port} p{i + 1}: workbook {got} vs oracle {want[i]}")
    return problems


def shard_of(items: list, index: int, count: int) -> list:
    """Every `count`-th item from `index`: the `count` shards together hold each item exactly once."""
    return [x for i, x in enumerate(items) if i % count == index]


def corruption_audit(
    comp: Composed,
    period: int = 2,
    shard: tuple[int, int] | None = None,
    modules: set[str] | None = None,
) -> list[str]:
    """Corrupt each module output row in one period, by +25%; some Summary check must turn non-zero.
    (Each module's own acceptance audits every cell of its rows; this one proves the wiring is covered.)
    `shard=(index, count)` audits one share of the output rows (the shards together are every row, once) and
    `modules` only those modules' rows; CI runs the shards in parallel jobs."""
    problems: list[str] = []
    ws_name = comp.calcs.ws.title
    col = comp.wb.period_cols[period - 1]
    rows = comp.calcs.rows()
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        comp.save(path)
        base = evaluate_cells(path, (ws_name,))
        check_cells = [
            f"{ws_name.upper()}'!{c}{r}"
            for _m, _l, r in comp.check_rows
            for c in comp.wb.period_cols
        ]
        clean_red = {k for k in check_cells if abs(base.get(k, 0.0)) > 1e-6}
        targets = [
            (name, port)
            for name, outs in comp.expected.items()
            if modules is None or name in modules
            for port in outs
        ]
        if shard is not None:
            targets = shard_of(targets, *shard)
        for name, port in targets:
            row = rows[f"{name}.{port}"]
            w = openpyxl.load_workbook(path)
            cell = w[ws_name][f"{col}{row}"]
            v = base.get(f"{ws_name.upper()}'!{col}{row}", 1.0)
            cell.value = f"=({cell.value[1:]})+{0.25 * max(abs(v), 1.0)!r}"
            bad = os.path.join(d, "bad.xlsx")
            w.save(bad)
            got = evaluate_cells(bad, (ws_name,))
            red = {k for k in check_cells if abs(got.get(k, 0.0)) > 1e-6} - clean_red
            if not red:
                problems.append(f"{name}.{port}: corrupted in period {period}, no check fired")
    return problems


def inputs_drive_outputs(
    comp: Composed,
    external: dict,
    rel_tol: float = 1e-6,
    moves: dict | None = None,
    shard: tuple[int, int] | None = None,
) -> list[str]:
    """Change each external input in the built workbook (+10%, or to 1 if it was 0) and re-evaluate: every module's
    output rows must equal the chained oracle on the changed inputs. A wire that carries a number baked in at build
    time, or an input that reaches nothing it should, fails here. `moves` maps an input name to its own change
    (a function of the value) where the default would leave the model's domain, for example a spend profile
    given spend after construction."""
    problems: list[str] = []
    moves = moves or {}
    rows = comp.calcs.rows()
    cols = list(comp.wb.period_cols)
    bump = lambda v: v * 1.1 if v != 0 else 1.0  # noqa: E731
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "t.xlsx")
        comp.save(path)
        cells = list(comp.ext_cells.items())
        if shard is not None:
            cells = shard_of(cells, *shard)
        for name, (sheet, where) in cells:
            move = moves.get(name)
            if move is not None:
                new = move(external[name])
            elif isinstance(external[name], list):
                new = [bump(float(v)) for v in external[name]]
            else:
                new = bump(float(external[name]))
            moved = {**external, name: new}
            w = openpyxl.load_workbook(path)
            if sheet == "Inputs":
                w["Inputs"][where] = moved[name]
            else:
                for i, v in enumerate(moved[name]):
                    w["Feed"].cell(row=where, column=3 + i, value=float(v))
            changed = os.path.join(d, "moved.xlsx")
            w.save(changed)
            values = evaluate_cells(changed, ("Calcs",))
            want = chained_oracle(comp.steps, moved)
            for mod, outs in want.items():
                for port, series in outs.items():
                    for i, col in enumerate(cols):
                        got = values.get(f"CALCS'!{col}{rows[f'{mod}.{port}']}")
                        if not isinstance(got, float) or abs(got - series[i]) > rel_tol * max(
                            1.0, abs(series[i])
                        ):
                            problems.append(
                                f"{name} +10%: {mod}.{port} p{i + 1}: {got} vs {series[i]}"
                            )
                            break
    return problems
