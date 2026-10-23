"""Style profiles: how a model looks, kept apart from what it calculates.

Structure (FAST layout, calculation blocks, every formula) is fixed. A style profile is a named,
swappable set of choices about presentation only: number formats, unit labels, shading, rules,
fonts and colours. Two profiles given the same model produce the same values and the same
formulas, and differ only in formats and styling.

`DEFAULT` carries the house rules:

* negatives in brackets: ``(1,234)``;
* decimals only where a value needs them, to ``sig_digits`` significant digits below 1, and never a
  trailing dot (``smart``); whole values show no decimals;
* prices to 2 decimal places, always (``PRICE``);
* unit labels as ``Label (unit)``;
* totals ruled above and double-ruled below, headers ruled with no fill;
* a white sheet (no gridlines), and shading on hard-coded input cells only.

An organisation's own look is a profile that replaces these choices; it adds no structure.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, replace

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

# Format tokens a caller may pass for a row. A profile turns each into a real number format.
NUM = "__smart__"  # decimals picked from the value
PRICE = "__price__"  # always 2 decimals
ACCT0 = "#,##0;(#,##0)"  # whole numbers
ACCT2 = "#,##0.00;(#,##0.00)"  # 2 decimals
RATE = "0.######"
DATEF = "yyyy-mm-dd"

RIGHT = Alignment(horizontal="right")


class StyleError(ValueError):
    pass


@dataclass(frozen=True)
class StyleProfile:
    name: str
    negative: str = "brackets"  # "brackets" -> (1,234); "minus" -> -1,234
    sig_digits: int = 2  # significant digits shown for a value below 1
    max_decimals: int = 6
    rate: str = RATE
    datef: str = DATEF
    unit_pattern: str = "{label} ({unit})"
    accent: str = "1F2A44"  # titles, headers, section rules
    muted: str = "55617A"  # subtitles
    rule: str = "8894AC"  # thin rule colour
    input_fill: str | None = "E8F0FE"  # shade on a hard-coded input; None = no shading
    total_above: str | None = "thin"
    total_below: str | None = "double"
    gridlines: bool = False
    font_name: str | None = None

    def __post_init__(self) -> None:
        if self.negative not in ("brackets", "minus"):
            raise StyleError(f"{self.name}: negative must be 'brackets' or 'minus'")
        if "{label}" not in self.unit_pattern or "{unit}" not in self.unit_pattern:
            raise StyleError(f"{self.name}: unit_pattern needs {{label}} and {{unit}}")
        for side in (self.total_above, self.total_below):
            if side not in (None, "thin", "double"):
                raise StyleError(f"{self.name}: a total rule is None, 'thin' or 'double'")

    # -- numbers -------------------------------------------------------------------------------
    def grouped(self, decimals: int) -> str:
        body = "#,##0" + ("." + "0" * decimals if decimals else "")
        return f"{body};({body})" if self.negative == "brackets" else f"{body};-{body}"

    @property
    def acct0(self) -> str:
        return self.grouped(0)

    @property
    def acct2(self) -> str:
        return self.grouped(2)

    def decimals_for(self, v: float) -> int:
        """Decimals a value needs: none if whole to 2 dp, 2 for cents, more below 1."""
        v = float(v)
        if v == 0:
            return 0
        if abs(v) >= 1:
            return 0 if round(v, 2) == int(round(v, 2)) else 2
        d = min(math.floor(-math.log10(abs(v))) + self.sig_digits, self.max_decimals)
        while d > 2 and round(v, d - 1) == round(v, d):
            d -= 1
        return max(d, 2)

    def smart_num(self, v: float) -> str:
        return self.grouped(self.decimals_for(v))

    def fmt_for(self, value, fmt: str = NUM) -> str:
        """Turn a format token (or an explicit format) into this profile's number format."""
        if fmt == NUM:
            if isinstance(value, bool):
                return self.acct0
            return self.smart_num(value) if isinstance(value, int | float) else self.acct2
        return self.token(fmt)

    def token(self, fmt: str) -> str:
        """A token or one of the standard format strings, in this profile's form."""
        if fmt == PRICE or fmt == ACCT2:
            return self.acct2
        if fmt == ACCT0:
            return self.acct0
        if fmt == RATE:
            return self.rate
        if fmt == DATEF:
            return self.datef
        return fmt

    def fallback(self, fmt: str) -> str:
        """A number format for a formula cell whose value is not known: never the sentinel."""
        return self.acct2 if fmt == NUM else self.token(fmt)

    # -- labels --------------------------------------------------------------------------------
    def unit_label(self, label: str, unit: str) -> str:
        return self.unit_pattern.format(label=label.rstrip(), unit=unit)

    # -- cells and sheets ----------------------------------------------------------------------
    def _font(self, **kw) -> Font:
        if self.font_name:
            kw.setdefault("name", self.font_name)
        return Font(**kw)

    def _side(self, kind: str | None, color: str) -> Side | None:
        return Side(style=kind, color=color) if kind else None

    @property
    def total_top(self) -> Border:
        return Border(top=self._side(self.total_above, self.rule))

    @property
    def total_bottom(self) -> Border:
        return Border(top=self._side(self.total_below, self.accent))

    @property
    def hdr_border(self) -> Border:
        return Border(bottom=Side(style="thin", color=self.accent))

    def bold(self, cell, size=None, color=None) -> None:
        cell.font = self._font(bold=True, size=size, color=color)

    def section(self, cell, size=None) -> None:
        self.bold(cell, size=size, color=self.accent)

    def title(self, cell, size: int = 12) -> None:
        cell.font = self._font(bold=True, size=size, color=self.accent)

    def subtitle(self, cell) -> None:
        cell.font = self._font(italic=True, color=self.muted)

    def colhdr(self, cell, right: bool = True) -> None:
        cell.font = self._font(bold=True, color=self.accent)
        cell.border = self.hdr_border
        if right:
            cell.alignment = RIGHT

    def mark_input(self, cell, fmt: str | None = None) -> None:
        if self.input_fill:
            cell.fill = PatternFill("solid", fgColor=self.input_fill)
        if fmt:
            cell.number_format = fmt

    def white(self, ws: Worksheet) -> None:
        ws.sheet_view.showGridLines = self.gridlines

    def with_(self, **changes) -> StyleProfile:
        """A copy with some choices changed, under a new name."""
        return replace(self, **changes)


DEFAULT = StyleProfile("omd-default")

#: A second, plainer profile: minus signs, no shading hierarchy beyond a pale yellow input cell,
#: single total rules, gridlines on. It exists to show that style swaps without touching structure.
PLAIN = StyleProfile(
    "plain",
    negative="minus",
    unit_pattern="{label} [{unit}]",
    accent="000000",
    muted="666666",
    rule="999999",
    input_fill="FFF2CC",
    total_above="thin",
    total_below="thin",
    gridlines=True,
)

PROFILES: dict[str, StyleProfile] = {DEFAULT.name: DEFAULT, PLAIN.name: PLAIN}


def register(profile: StyleProfile) -> StyleProfile:
    if profile.name in PROFILES and PROFILES[profile.name] != profile:
        raise StyleError(f"a different profile named {profile.name!r} is already registered")
    PROFILES[profile.name] = profile
    return profile


def get_profile(style: str | StyleProfile | None) -> StyleProfile:
    """A profile by name or instance; None gives the default."""
    if style is None:
        return DEFAULT
    if isinstance(style, StyleProfile):
        return style
    try:
        return PROFILES[style]
    except KeyError:
        raise StyleError(f"no style profile {style!r}; known: {sorted(PROFILES)}") from None


def audit_style(
    path: str, profile: StyleProfile = DEFAULT, *, inputs_sheet: str = "Inputs"
) -> list[str]:
    """Check a built workbook against a profile. Returns one line per breach (empty = clean).

    * only hard-coded input cells (value column of the inputs sheet) carry the input shading;
    * gridlines are as the profile says;
    * a grouped number format follows the profile's negative style.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path)
    problems: list[str] = []
    want_fill = (profile.input_fill or "").upper()
    for ws in wb.worksheets:
        if bool(ws.sheet_view.showGridLines) != profile.gridlines:
            problems.append(
                f"{ws.title}: gridlines {'on' if ws.sheet_view.showGridLines else 'off'}"
            )
        for row in ws.iter_rows():
            for c in row:
                if c.fill is not None and c.fill.fill_type == "solid":
                    color = str(c.fill.fgColor.rgb or "").upper()
                    is_input = (
                        ws.title == inputs_sheet and c.column == 2 and color.endswith(want_fill)
                    )
                    if not is_input:
                        problems.append(f"{ws.title}!{c.coordinate}: shaded but not an input")
                nf = c.number_format or ""
                if "#,##0" in nf:
                    has_brackets = "(" in nf
                    if profile.negative == "brackets" and not has_brackets:
                        problems.append(
                            f"{ws.title}!{c.coordinate}: negatives not in brackets ({nf})"
                        )
                    if profile.negative == "minus" and has_brackets:
                        problems.append(f"{ws.title}!{c.coordinate}: bracketed negatives ({nf})")
    return problems
