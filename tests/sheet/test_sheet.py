"""Tests for omd.sheet: the synthetic reference model proves the builder is generic, its
live formulas reproduce the Python source-of-truth, house formatting is applied, and the build
is content-deterministic.
"""

from __future__ import annotations

import pytest

from omd.sheet import formatting as F
from omd.sheet.example import build_reference, source_of_truth
from omd.sheet.verify import verify


# ---- smart number format (the house rule) ----------------------------------------------------
@pytest.mark.parametrize(
    "value,expected",
    [
        (1200.0, F.ACCT0),  # whole -> no decimals
        (800.5, F.ACCT2),  # fractional -> two decimals
        (-450.75, F.ACCT2),
        (-300.0, F.ACCT0),
        (0.0, F.ACCT0),
    ],
)
def test_smart_num(value, expected):
    assert F.smart_num(value) == expected


def test_fmt_for_resolves_sentinel():
    assert F.fmt_for(10.0, F.NUM) == F.ACCT0
    assert F.fmt_for(10.5, F.NUM) == F.ACCT2
    assert F.fmt_for(10.5, F.RATE) == F.RATE  # explicit format passes through


# ---- the reference model builds and verifies -------------------------------------------------
def test_reference_verifies(tmp_path):
    """The generated workbook's live formulas reproduce the Python source-of-truth."""
    pytest.importorskip("formulas")
    path = tmp_path / "reference.xlsx"
    _wb, calc, sot = build_reference(str(path))
    verify(str(path), calc, sot, tol=0.01)


def test_reference_sheets_and_formatting(tmp_path):
    import openpyxl

    path = tmp_path / "reference.xlsx"
    build_reference(str(path))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Inputs", "Ledger", "Calcs", "Accounts"]

    inputs = wb["Inputs"]
    # opening cash 5000 is a shaded input, whole -> ACCT0 (no trailing dot)
    b5 = inputs["B5"]
    assert b5.value == 5000.0
    assert b5.fill.fgColor.rgb.endswith("E8F0FE")  # INPUT_FILL
    assert b5.number_format == F.ACCT0
    assert inputs.sheet_view.showGridLines is False

    calcs = wb["Calcs"]
    # the Costs line carries fractional values in Feb -> smart format picks 2dp on a FORMULA cell
    labels = {
        str(calcs.cell(row=r, column=1).value).strip(): r for r in range(1, calcs.max_row + 1)
    }
    cost_row = labels["Costs"]
    feb = calcs.cell(row=cost_row, column=4)  # column D = Feb
    assert str(feb.value).startswith("=SUMIFS(")
    assert feb.number_format == F.ACCT2  # -450.75 has cents
    jan = calcs.cell(row=cost_row, column=3)  # -300.00 is whole
    assert jan.number_format == F.ACCT0

    accounts = wb["Accounts"]
    # statement cells link back to Calcs and render as whole numbers
    assert any(
        str(c.value).startswith("=Calcs!") for row in accounts.iter_rows() for c in row if c.value
    )


def test_build_is_content_deterministic(tmp_path):
    a = build_reference(str(tmp_path / "a.xlsx"))[0].cells()
    b = build_reference(str(tmp_path / "b.xlsx"))[0].cells()
    assert a == b, "workbook content is not deterministic"


def test_source_of_truth_math():
    sot = source_of_truth()
    # Jan: revenue 1200 + 800.5 = 2000.5; cost -300; profit 1700.5
    assert sot["rev"][0] == pytest.approx(2000.5)
    assert sot["profit"][0] == pytest.approx(1700.5)
    # cash Jan = 5000 + profit(1700.5) + draw(5000) = 11700.5
    assert sot["cash"][0] == pytest.approx(11700.5)
    # loan Jan = 10000 + 5000 = 15000; Mar adds 2500 -> 17500
    assert sot["loan"][0] == pytest.approx(15000.0)
    assert sot["loan"][2] == pytest.approx(17500.0)


def test_verify_goes_red_on_corrupted_oracle(tmp_path):
    """Negative control: a drifted source-of-truth must fail verify, naming the line."""
    path = str(tmp_path / "ref.xlsx")
    _wb, calc, sot = build_reference(path)
    verify(path, calc, sot, tol=0.01)  # control is meaningful only if the clean case is green
    bad = {k: list(v) for k, v in sot.items()}
    key = next(iter(bad))
    bad[key][0] += 1000.0
    with pytest.raises(AssertionError, match=key):
        verify(path, calc, bad, tol=0.01)
