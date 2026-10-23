"""The width linter: equal value columns, a notes column that does not wrap."""

import openpyxl

from omd.sheet import formatting as F


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    return ws


def test_value_columns_come_out_equal():
    ws = _sheet()
    ws["C2"], ws["D2"], ws["E2"] = 1.0, 22_222_222.0, 333.0
    for c in "CDE":
        ws[f"{c}2"].number_format = "#,##0"
    F.lint_widths(ws, value_cols="CDE")
    w = {c: ws.column_dimensions[c].width for c in "CDE"}
    assert len(set(w.values())) == 1, f"value columns differ: {w}"


def test_a_note_in_a_value_column_is_REPORTED_not_widened_around():
    """The bug this linter was written for: a note written into the Total column."""
    ws = _sheet()
    ws["E2"] = "Zone A only, converted at the NGL density — NGL is sold by volume, not weight"
    out = F.lint_widths(ws, value_cols="CDE", value_max=24)
    assert any("wrong column" in f for f in out), out
    assert ws.column_dimensions["E"].width <= 24, "a stray note must not widen a value column"


def test_note_column_is_sized_not_to_wrap():
    ws = _sheet()
    ws["F2"] = "A" * 60
    F.lint_widths(ws, value_cols="CDE", note_col="F", note_max=120)
    assert ws.column_dimensions["F"].width >= 60


def test_a_note_past_the_cap_is_reported():
    ws = _sheet()
    ws["F2"] = "A" * 200
    out = F.lint_widths(ws, value_cols="CDE", note_col="F", note_max=120)
    assert any("will wrap" in f for f in out), out
    assert ws.column_dimensions["F"].width == 120


def test_a_clean_sheet_reports_nothing():
    ws = _sheet()
    ws["C2"], ws["D2"], ws["E2"] = 1.0, 2.0, 3.0
    ws["F2"] = "A short note"
    assert F.lint_widths(ws, value_cols="CDE", note_col="F") == []


def test_footer_rows_can_be_skipped():
    """A provenance footer is not a table: its prose must not read as a mis-placed note."""
    ws = _sheet()
    ws["C90"] = "build_workbook.py (omd spec module) — regenerate, never hand-edit"
    assert F.lint_widths(ws, value_cols="CDE") != []  # seen by default
    assert F.lint_widths(ws, value_cols="CDE", skip_rows=range(85, 100)) == []
