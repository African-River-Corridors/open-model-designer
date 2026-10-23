"""Nothing a module writes may depend on what sits right of the last period: stray numbers there change no output."""

import os
import tempfile

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string, get_column_letter

from omd import acceptance
from omd.harness._eval import evaluate_cells
from omd.registry import discover

MODULES = discover("omd.modules")


@pytest.mark.parametrize("name", sorted(MODULES))
def test_a_number_right_of_the_last_period_changes_no_output(name):
    m = MODULES[name]
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "m.xlsx")
        acceptance._host(m, m.example, path)
        base = evaluate_cells(path, ("Calcs",))
        wb = openpyxl.load_workbook(path)
        beyond = column_index_from_string(acceptance.PERIOD_COLS[-1]) + 1
        for ws in wb.worksheets:
            for r in range(1, ws.max_row + 1):
                for k in range(3):  # the three columns right of the last period, on every sheet
                    ws.cell(row=r, column=beyond + k, value=12345.0)
        noisy = os.path.join(d, "noisy.xlsx")
        wb.save(noisy)
        got = evaluate_cells(noisy, ("Calcs",))
        moved = [
            k
            for k, v in base.items()
            if isinstance(v, float)
            and k.split("!")[1].rstrip("0123456789") in acceptance.PERIOD_COLS
            and abs(got.get(k, float("nan")) - v) > 1e-9
        ]
        assert moved == [], (
            f"{name}: cells moved when numbers sat right of the last period: {moved[:5]}"
        )
        assert get_column_letter(beyond) == "G"
