"""T6: structure is fixed, style is a swappable profile."""

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from omd.sheet import formatting as F
from omd.sheet.example import build_reference
from omd.sheet.verify import verify
from omd.style import DEFAULT, PLAIN, PRICE, StyleError, StyleProfile, audit_style, get_profile


def cells(path):
    """Every cell's (value or formula) by sheet and coordinate."""
    wb = openpyxl.load_workbook(path)
    return {
        (ws.title, c.coordinate): c.value for ws in wb.worksheets for r in ws.iter_rows() for c in r
    }


def formats(path):
    wb = openpyxl.load_workbook(path)
    return {
        (ws.title, c.coordinate): c.number_format
        for ws in wb.worksheets
        for r in ws.iter_rows()
        for c in r
        if c.value is not None
    }


@pytest.fixture(scope="module")
def two(tmp_path_factory):
    d = tmp_path_factory.mktemp("s")
    out = {}
    for prof in (DEFAULT, PLAIN):
        path = str(d / f"{prof.name}.xlsx")
        _wb, calc, sot = build_reference(path, style=prof)
        out[prof.name] = (path, calc, sot)
    return out


# --- the gate: same model, two profiles ---------------------------------------------------------
def test_two_profiles_render_identical_values_and_formulas(two):
    a, b = two["omd-default"][0], two["plain"][0]
    assert cells(a) == cells(b)


def test_two_profiles_differ_in_formats(two):
    fa, fb = formats(two["omd-default"][0]), formats(two["plain"][0])
    assert fa.keys() == fb.keys()
    differing = [k for k in fa if fa[k] != fb[k]]
    assert len(differing) > 20
    assert any("(" in fa[k] for k in differing) and not any("(" in fb[k] for k in differing)


def test_both_profiles_still_verify_against_the_oracle(two):
    for path, calc, sot in two.values():
        verify(path, calc, sot, tol=0.01)


def test_the_profiles_differ_in_look_too(two):
    wa = openpyxl.load_workbook(two["omd-default"][0])
    wp = openpyxl.load_workbook(two["plain"][0])
    assert wa["Calcs"].sheet_view.showGridLines is False
    assert wp["Calcs"].sheet_view.showGridLines is True
    ia = wa["Inputs"]["B5"].fill.fgColor.rgb
    ip = wp["Inputs"]["B5"].fill.fgColor.rgb
    assert ia.endswith("E8F0FE") and ip.endswith("FFF2CC")


def test_the_default_profile_is_what_the_builder_uses_without_a_style(tmp_path):
    p1, p2 = str(tmp_path / "a.xlsx"), str(tmp_path / "b.xlsx")
    build_reference(p1)
    build_reference(p2, style="omd-default")
    assert formats(p1) == formats(p2)


# --- the house rules, in the default profile ------------------------------------------------------
def test_negatives_in_brackets():
    assert DEFAULT.acct0 == "#,##0;(#,##0)" and DEFAULT.acct2 == "#,##0.00;(#,##0.00)"
    assert PLAIN.acct0 == "#,##0;-#,##0"


@pytest.mark.parametrize(
    "value, decimals",
    [(1000, 0), (1000.004, 0), (1234.5, 2), (0.35, 2), (0.05, 2), (0.0042, 4), (0.00042, 5),
     (-0.0042, 4), (0, 0), (0.0000045, 6)],
)  # fmt: skip
def test_significant_digit_decimals(value, decimals):
    assert DEFAULT.decimals_for(value) == decimals
    fmt = DEFAULT.smart_num(value)
    assert fmt.count("0") >= 1 and not fmt.split(";")[0].endswith(".")  # never a trailing dot


def test_prices_are_always_two_decimals():
    assert DEFAULT.fmt_for(12, PRICE) == DEFAULT.acct2
    assert DEFAULT.fmt_for(12, F.NUM) == DEFAULT.acct0  # the same value as a plain number
    assert DEFAULT.fmt_for(12.0, F.PRICE) == DEFAULT.acct2


def test_unit_labels():
    assert DEFAULT.unit_label("Unit price  ", "USD/tonne") == "Unit price (USD/tonne)"
    assert PLAIN.unit_label("Unit price", "USD/tonne") == "Unit price [USD/tonne]"


def test_standard_format_strings_follow_the_profile():
    assert PLAIN.fmt_for(1, F.ACCT0) == "#,##0;-#,##0" and PLAIN.fmt_for(1, F.ACCT2) == PLAIN.acct2
    assert PLAIN.fmt_for(1, "0.0%") == "0.0%"  # an explicit format passes through


def test_total_rules_above_and_below(tmp_path):
    from omd.harness import fastblock
    from omd.harness.fastblock import Block

    for prof, below in ((DEFAULT, "double"), (PLAIN, "thin")):
        from omd.sheet.workbook import Workbook

        fastblock.TREE.__init__()  # the driver tree is module-level state (see T4 notes)
        w = Workbook("T", ["a", "b"], opening=True, style=prof)
        ins = w.inputs("D", {"x": ("X", 2.0)})
        c = w.calcs()
        b = Block(c, "B", registry={})
        b.link("x", ins["x"], "X ($)", values=[2.0, 2.0])
        b.calc("y", "Y ($)", "{x}*{x}", values=[4.0, 4.0])
        rr = b.total("t", "Total ($)", ["y"])
        path = str(tmp_path / f"{prof.name}.xlsx")
        w.save(path)
        ws = openpyxl.load_workbook(path)["Calcs"]
        assert ws[f"C{rr}"].border.top.style == "thin"
        assert ws[f"C{rr + 1}"].border.top.style == below
        assert ws[f"A{rr}"].font.bold


# --- the audit, with negative controls ------------------------------------------------------------
def test_a_clean_workbook_passes_its_own_profile(two):
    assert audit_style(two["omd-default"][0], DEFAULT) == []
    assert audit_style(two["plain"][0], PLAIN) == []


def test_the_audit_goes_red_when_a_calc_cell_is_shaded(two, tmp_path):
    """Negative control: shade one Calcs cell; only inputs may be shaded."""
    path = str(tmp_path / "shaded.xlsx")
    w = openpyxl.load_workbook(two["omd-default"][0])
    w["Calcs"]["C6"].fill = PatternFill("solid", fgColor="E8F0FE")
    w.save(path)
    assert audit_style(path, DEFAULT) == ["Calcs!C6: shaded but not an input"]


def test_the_audit_goes_red_on_the_wrong_profile(two):
    """The plain workbook is not a default-profile workbook: gridlines, fills and minus signs."""
    problems = audit_style(two["plain"][0], DEFAULT)
    assert any("gridlines" in p for p in problems)
    assert any("shaded but not an input" in p for p in problems)
    assert any("negatives not in brackets" in p for p in problems)


def test_the_audit_goes_red_on_a_bracketless_format(two, tmp_path):
    path = str(tmp_path / "nobracket.xlsx")
    w = openpyxl.load_workbook(two["omd-default"][0])
    w["Calcs"]["C6"].number_format = "#,##0.00"
    w.save(path)
    assert audit_style(path, DEFAULT) == ["Calcs!C6: negatives not in brackets (#,##0.00)"]


# --- profiles ------------------------------------------------------------------------------------
def test_profiles_by_name_and_instance():
    assert get_profile(None) is DEFAULT and get_profile("plain") is PLAIN
    mine = DEFAULT.with_(name="mine", accent="123456")
    assert get_profile(mine) is mine and mine.name == "mine"
    with pytest.raises(StyleError, match="no style profile 'nope'"):
        get_profile("nope")


@pytest.mark.parametrize(
    "bad",
    [
        {"negative": "red"},
        {"unit_pattern": "{label}"},
        {"total_above": "triple"},
    ],
)
def test_a_bad_profile_is_refused(bad):
    with pytest.raises(StyleError):
        StyleProfile("x", **bad)
