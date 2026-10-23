"""A small model built with the harness, plus the helpers the harness tests share."""

from __future__ import annotations

import shutil

import openpyxl

from omd.harness import fastblock
from omd.harness.fastblock import Block
from omd.sheet.workbook import Workbook

PERIODS = ["P1", "P2", "P3", "P4"]
VOL = [1.0, 2.0, 3.0, 4.0]
PRICE = 10.0
REV = [PRICE * v for v in VOL]
RATE = 0.2
TAX = [r * RATE for r in REV]
LAST = "F"  # periods sit in C..F: A = label, B = Open


def reset_registry():
    fastblock.TREE.__init__()
    fastblock.BLOCK_VALUES.clear()


def build(path: str) -> str:
    """Revenue and tax blocks, a re-derivation check on revenue and NONE on tax."""
    reset_registry()
    wb = Workbook("Toy", PERIODS, opening=True)
    ins = wb.inputs(
        "Drivers",
        {
            "price": ("Unit price", PRICE),
            "tax_rate": ("Tax rate", RATE),
            "dead": ("Dead input", 5.0),
        },
    )
    sched = wb._wb.create_sheet("Schedules")
    sched["A2"] = "Volume"
    for i, v in enumerate(VOL):
        sched.cell(row=2, column=3 + i, value=v)
    calcs = wb.calcs()
    reg: dict = {}
    b = Block(calcs, "Revenue", registry=reg)
    b.link("price", ins["price"], "Unit price ($)", values=[PRICE] * 4)
    b.link("vol", "Schedules!{col}2", "Volume (units)", values=VOL)
    b.calc("rev", "Revenue ($)", "{price}*{vol}", key="rev")
    b2 = Block(calcs, "Tax", registry=reg)
    b2.link("rev", "rev", "Revenue ($)", values=REV)
    b2.link("rate", ins["tax_rate"], "Tax rate (%)", values=[RATE] * 4)
    b2.calc("tax", "Tax ($)", "{rev}*{rate}", key="tax")
    calcs.section("Checks")
    calcs.row(
        "chk_rev",
        "  Check — revenue re-derived ($)",
        lambda c: f"={c.here('rev')}-{ins['price']}*Schedules!{c.col}2",
        values=[0.0] * 4,
    )
    rows = calcs.rows()
    chk = rows["chk_rev"]
    tax = rows["tax"]

    summ = wb._wb.create_sheet("Summary")
    summ["B1"], summ["C1"] = "check", "status"
    total = "+".join(f"ABS(Calcs!{c}{chk})" for c in "CDEF")
    summ["B2"] = "Revenue re-derived"
    summ["C2"] = f'=IF({total}<1,"OK","FAIL")'

    checks = wb._wb.create_sheet("Checks")
    checks["B1"] = "Validity checks"
    for r, (lab, f) in enumerate(
        [("Revenue re-derived", summ["C2"].value), ("Tax is present", '="OK"')], start=3
    ):
        checks.cell(row=r, column=2, value=lab)
        checks.cell(row=r, column=3, value=f.replace("Calcs!", "Calcs!"))
        checks.cell(row=r, column=4, value="error")

    ann = wb._wb.create_sheet("Annual Summary")
    ann["B1"] = "Check"
    ann["B2"] = "Revenue re-derived"
    ann["B3"] = "Tax is not negative (≥ 0)"
    for i, c in enumerate("CDEF"):
        ann.cell(row=2, column=3 + i, value=f"=Calcs!{c}{chk}")
        ann.cell(row=3, column=3 + i, value=f"=Calcs!{c}{tax}")
    wb.save(path)
    return path


def corrupt(src: str, dst: str, sheet: str, ref: str, factor: float = 1.25) -> str:
    """Scale one formula cell by `factor` in a copy of the workbook."""
    shutil.copy(src, dst)
    w = openpyxl.load_workbook(dst)
    old = w[sheet][ref].value
    w[sheet][ref] = "=(" + str(old)[1:] + f")*{factor:g}"
    w.save(dst)
    return dst


def row_of(path: str, label: str, sheet: str = "Calcs") -> int:
    ws = openpyxl.load_workbook(path)[sheet]
    return next(
        r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=1).value).strip() == label
    )
